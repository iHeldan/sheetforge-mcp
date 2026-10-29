"""Bounded, declarative native-table report recipes."""

from __future__ import annotations

from contextlib import contextmanager
import hashlib
import hmac
import json
import logging
import math
from pathlib import Path
import re
import shutil
import tempfile
from typing import Any

from openpyxl.cell.cell import ERROR_CODES
from openpyxl.compat import safe_string
from openpyxl.utils import get_column_letter

from . import changeset, query, workbook
from .exceptions import PreconditionFailedError, ResponseTooLargeError, ValidationError
from .tables import _find_table, _table_data_bounds
from .workbook import safe_workbook

RECIPE_ENGINE_VERSION = 1
MAX_RECIPE_BYTES = 32_000
MAX_SOURCE_CELLS = 50_000
MAX_OUTPUT_CELLS = 10_000
MAX_FIELDS = 64
MAX_SUMMARY_BYTES = 8_000
NUMERIC_REL_TOL = 1e-12
NUMERIC_ABS_TOL = 1e-9
logger = logging.getLogger(__name__)


@contextmanager
def _recipe_workspace():
    path = Path(tempfile.mkdtemp(prefix="sheetforge-recipe-"))
    try:
        yield path
    finally:
        # Cleanup must not turn an already committed transaction into a reported
        # failure. The ChangeSet owns persistence and rollback, not this copy.
        try:
            shutil.rmtree(path)
        except OSError:
            logger.warning("Could not remove recipe temporary workspace: %s", path)


def _object(
    value: Any, *, allowed: set[str], required: set[str], location: str
) -> dict:
    if not isinstance(value, dict):
        raise ValidationError(f"{location} must be an object")
    unknown, missing = set(value) - allowed, required - set(value)
    if unknown or missing:
        raise ValidationError(
            f"{location}: unknown fields {sorted(unknown)}, missing fields {sorted(missing)}"
        )
    return value


def _name(value: Any, location: str) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > 128:
        raise ValidationError(
            f"{location} must be a nonblank string of at most 128 characters"
        )
    return value


def _field_list(value: Any, location: str) -> list[str]:
    if not isinstance(value, list) or len(value) > MAX_FIELDS:
        raise ValidationError(
            f"{location} must be a list of at most {MAX_FIELDS} field names"
        )
    result = [_name(item, location) for item in value]
    if len(set(result)) != len(result):
        raise ValidationError(f"{location} contains duplicate field names")
    return result


def _normalize_recipe(recipe: Any) -> dict:
    recipe = changeset._json_copy(recipe, location="recipe")
    if len(changeset._canonical_json(recipe).encode("utf-8")) > MAX_RECIPE_BYTES:
        raise ValidationError(f"recipe exceeds {MAX_RECIPE_BYTES} bytes")
    _object(
        recipe,
        allowed={
            "version",
            "source",
            "filters",
            "group_by",
            "metrics",
            "output",
            "checks",
        },
        required={"version", "source", "group_by", "metrics", "output"},
        location="recipe",
    )
    if type(recipe["version"]) is not int or recipe["version"] != 1:
        raise ValidationError("recipe.version must be 1")
    source = _object(
        recipe["source"],
        allowed={"table_name", "sheet_name"},
        required={"table_name"},
        location="source",
    )
    for key, value in source.items():
        _name(value, f"source.{key}")
    output = _object(
        recipe["output"],
        allowed={"sheet_name", "range_ref", "allow_overwrite"},
        required={"sheet_name", "range_ref"},
        location="output",
    )
    sheet_name = _name(output["sheet_name"], "output.sheet_name")
    if len(sheet_name) > 31 or any(char in sheet_name for char in "\\/*?:[]"):
        raise ValidationError("output.sheet_name must be a valid Excel worksheet name")
    output["range_ref"] = changeset._canonical_range_ref(
        output["range_ref"], location="output.range_ref"
    )
    if (
        changeset._range_cell_count(output["range_ref"], location="output.range_ref")
        > MAX_OUTPUT_CELLS
    ):
        raise ValidationError(f"output range exceeds {MAX_OUTPUT_CELLS} cells")
    output.setdefault("allow_overwrite", False)
    if type(output["allow_overwrite"]) is not bool:
        raise ValidationError("output.allow_overwrite must be a boolean")
    recipe["group_by"] = _field_list(recipe["group_by"], "group_by")
    metrics = recipe["metrics"]
    if not isinstance(metrics, list) or not 1 <= len(metrics) <= MAX_FIELDS:
        raise ValidationError(f"metrics must contain 1 to {MAX_FIELDS} entries")
    for metric in metrics:
        _object(
            metric,
            allowed={"op", "field", "as"},
            required={"op", "as"},
            location="metric",
        )
        _name(metric["op"], "metric.op")
        if metric["op"] not in query.AGGREGATE_OPERATORS:
            raise ValidationError("metric.op is not a supported aggregate operator")
        _name(metric["as"], "metric.as")
        if metric["op"] == "count":
            if "field" in metric:
                raise ValidationError("count counts rows and must not specify field")
        else:
            _name(metric.get("field"), "metric.field")
    recipe.setdefault("filters", [])
    if not isinstance(recipe["filters"], list) or len(recipe["filters"]) > 50:
        raise ValidationError("filters must be a list of at most 50 objects")
    for item in recipe["filters"]:
        _object(
            item,
            allowed={"field", "op", "value", "values", "case_sensitive"},
            required={"field", "op"},
            location="filter",
        )
        _name(item["field"], "filter.field")
        _name(item["op"], "filter.op")
        if "case_sensitive" in item and type(item["case_sensitive"]) is not bool:
            raise ValidationError("filter.case_sensitive must be a boolean")
        if "value" in item and "values" in item:
            raise ValidationError("filter must not specify both value and values")
        op = query.FILTER_OPERATOR_ALIASES.get(item["op"], item["op"])
        item["op"] = op
        if op in {"is_blank", "not_blank"} and ({"value", "values"} & set(item)):
            raise ValidationError("Blank checks must not specify value or values")
        if "values" in item and op not in {"in", "not_in"}:
            raise ValidationError("Only in/not_in filters accept values")
        for key in ("value", "values"):
            if key in item:
                values = item[key] if key == "values" else [item[key]]
                if not isinstance(values, list) or any(
                    isinstance(v, (dict, list)) for v in values
                ):
                    raise ValidationError("filter values must be JSON scalars")
    recipe.setdefault("checks", [])
    if not isinstance(recipe["checks"], list) or len(recipe["checks"]) > 20:
        raise ValidationError("checks must be a list of at most 20 objects")
    for check in recipe["checks"]:
        if not isinstance(check, dict):
            raise ValidationError("check must be an object")
        kind = check.get("type")
        if kind == "unique_key":
            _object(
                check,
                allowed={"type", "fields"},
                required={"type", "fields"},
                location="check",
            )
            if not _field_list(check["fields"], "check.fields"):
                raise ValidationError("unique_key.fields must not be empty")
        elif kind == "row_count":
            _object(
                check,
                allowed={"type", "min", "max"},
                required={"type", "min", "max"},
                location="check",
            )
            if any(
                type(check[key]) is not int or check[key] < 0 for key in ("min", "max")
            ):
                raise ValidationError("row_count bounds must be nonnegative integers")
            if check["min"] > check["max"]:
                raise ValidationError("row_count.min must not exceed max")
        elif kind == "sum_matches":
            _object(
                check,
                allowed={"type", "source_field", "output_field"},
                required={"type", "source_field", "output_field"},
                location="check",
            )
            _name(check["source_field"], "check.source_field")
            _name(check["output_field"], "check.output_field")
        else:
            raise ValidationError(
                "check.type must be unique_key, row_count, or sum_matches"
            )
    return recipe


def _source_fields(recipe: dict) -> list[str]:
    fields = list(recipe["group_by"])
    fields += [item["field"] for item in recipe["metrics"] if "field" in item]
    fields += [item["field"] for item in recipe["filters"]]
    for check in recipe["checks"]:
        if check["type"] == "unique_key":
            fields += check["fields"]
        elif check["type"] == "sum_matches":
            fields.append(check["source_field"])
    return list(dict.fromkeys(fields))


def _scalar(value: Any, *, location: str) -> Any:
    if value is None or type(value) in (str, bool, int):
        return value
    if type(value) is float and math.isfinite(value):
        return value
    raise ValidationError(
        f"{location} must contain finite JSON scalar values; dates and formulas are not supported"
    )


def _load_table(wb: Any, recipe: dict) -> dict:
    sheet, ws, table = _find_table(wb, **recipe["source"])
    changeset._range_bounds(table.ref, location="source table range")
    c1, r1, c2, r2, data_start, data_end = _table_data_bounds(table)
    if table.headerRowCount != 1:
        raise ValidationError("Recipes require a native table with one header row")
    if int(table.totalsRowCount or 0) not in (0, 1) or r2 - r1 < int(
        table.totalsRowCount or 0
    ):
        raise ValidationError("Source table has inconsistent totals-row bounds")
    if (c2 - c1 + 1) * (r2 - r1 + 1) > MAX_SOURCE_CELLS or c2 - c1 + 1 > MAX_FIELDS:
        raise ValidationError(
            f"Source table exceeds {MAX_SOURCE_CELLS} cells or {MAX_FIELDS} columns"
        )
    headers = [ws.cell(r1, col).value for col in range(c1, c2 + 1)]
    for header in headers:
        _name(header, "source header")
    if len({header.strip().casefold() for header in headers}) != len(headers):
        raise ValidationError("Source table has ambiguous duplicate headers")
    fields = _source_fields(recipe)
    missing = sorted(set(fields) - set(headers))
    if missing:
        raise ValidationError(
            f"Source fields not found (exact header names required): {missing}"
        )
    indexes = [headers.index(field) for field in fields]
    rows = []
    for row_number in range(data_start, data_end + 1):
        values = [None] * len(headers)
        for index in indexes:
            cell = ws.cell(row_number, c1 + index)
            if cell.data_type in {"f", "e"}:
                raise ValidationError(
                    f"Source {sheet}!{cell.coordinate} is a formula or Excel error; recipes require literal inputs"
                )
            values[index] = _scalar(
                cell.value, location=f"Source {sheet}!{cell.coordinate}"
            )
        rows.append(values)
    # Prevent Python's equality rules from merging True with 1, or filters from
    # silently dropping mixed numeric/text values in a supposedly numeric field.
    for index in indexes:
        kinds = {
            (
                "number"
                if type(row[index]) in (int, float)
                else type(row[index]).__name__
            )
            for row in rows
            if row[index] is not None
        }
        if len(kinds) > 1:
            raise ValidationError(
                f"Source field '{headers[index]}' contains mixed value types"
            )
    for item in recipe["filters"]:
        index = headers.index(item["field"])
        nonblank = next((row[index] for row in rows if row[index] is not None), None)
        if (
            item["op"] in {"contains", "starts_with", "ends_with"}
            and nonblank is not None
            and not isinstance(nonblank, str)
        ):
            raise ValidationError(
                f"String filter requires a text field: '{item['field']}'"
            )
        operands = item.get("values", [item["value"]] if "value" in item else [])
        for value in operands:
            if (
                nonblank is not None
                and value is not None
                and _kind(nonblank) != _kind(value)
            ):
                raise ValidationError(
                    f"Filter operand type does not match field '{item['field']}'"
                )
    return {
        "sheet_name": sheet,
        "table_name": table.displayName,
        "range_ref": table.ref,
        "headers": headers,
        "rows": rows,
        "data_start_row": data_start,
        "bindings": {
            field: f"{get_column_letter(c1 + headers.index(field))}{r1}"
            for field in fields
        },
        "totals_rows_excluded": int(table.totalsRowCount or 0),
    }


def _aggregate(source: dict, recipe: dict) -> dict:
    result = query._aggregate_dataset(
        headers=source["headers"],
        rows=source["rows"],
        target_kind="excel_table",
        sheet_name=source["sheet_name"],
        table_name=source["table_name"],
        auto_selected_sheet=False,
        filters=recipe["filters"],
        group_by=recipe["group_by"],
        metrics=recipe["metrics"],
        sort_by=None,
        sort_desc=False,
        limit=None,
        row_mode="arrays",
        infer_schema=False,
    )
    group_count = len(recipe["group_by"])
    # Assertions must describe the numbers the writer can actually persist, not
    # binary-float tails such as 0.1 + 0.2 == 0.30000000000000004. Do not weaken
    # ChangeSet's exact postconditions or quantize source/group identity values.
    for row in result["rows"]:
        for index in range(group_count, len(row)):
            if type(row[index]) is float:
                _scalar(row[index], location="Aggregate output")
                row[index] = float(safe_string(row[index]))
    result["rows"].sort(key=lambda row: changeset._canonical_json(row[:group_count]))
    return result


def _kind(value: Any) -> str:
    return "number" if type(value) in (int, float) else type(value).__name__


def _equal_result(left: list, right: list) -> bool:
    if len(left) != len(right):
        return False
    for a, b in zip(left, right):
        if len(a) != len(b):
            return False
        for x, y in zip(a, b):
            if _kind(x) != _kind(y):
                return False
            if _kind(x) == "number":
                if type(x) is int and type(y) is int:
                    if x != y:
                        return False
                elif not math.isclose(
                    x, y, rel_tol=NUMERIC_REL_TOL, abs_tol=NUMERIC_ABS_TOL
                ):
                    return False
            elif x != y:
                return False
    return True


def _sum_values(rows: list, index: int) -> int | float:
    values = [row[index] for row in rows if row[index] is not None]
    if any(_kind(value) != "number" for value in values):
        raise ValidationError("sum_matches requires numeric fields")
    if all(type(value) is int for value in values):
        return sum(values)
    try:
        return math.fsum(values)
    except (OverflowError, ValueError) as exc:
        raise ValidationError(
            "sum_matches exceeds the supported finite numeric range"
        ) from exc


def _verify_recipe(source: dict, recipe: dict, result: dict, sample_limit: int) -> dict:
    rows, headers = source["rows"], source["headers"]
    schema = query._build_schema(headers, rows)
    filters = query._normalize_filters(recipe["filters"], headers, schema)
    matched = [
        (source["data_start_row"] + index, row)
        for index, row in enumerate(rows)
        if query._apply_filters([row], filters)
    ]
    checks = []
    for check in recipe["checks"]:
        kind = check["type"]
        if kind == "unique_key":
            indexes = [headers.index(field) for field in check["fields"]]
            keys = [tuple(row[index] for index in indexes) for row in rows]
            blank = sum(any(value in (None, "") for value in key) for key in keys)
            duplicates = len(keys) - len(set(keys))
            checks.append(
                {
                    **check,
                    "passed": not blank and not duplicates,
                    "scope": "all_source_rows",
                    "blank_keys": blank,
                    "duplicate_rows": duplicates,
                }
            )
        elif kind == "row_count":
            count = len(result["rows"])
            checks.append(
                {
                    **check,
                    "passed": check["min"] <= count <= check["max"],
                    "actual": count,
                }
            )
        else:
            if check["output_field"] not in result["headers"]:
                raise ValidationError(
                    f"sum_matches output field '{check['output_field']}' not found"
                )
            expected = _sum_values(
                [row for _, row in matched], headers.index(check["source_field"])
            )
            actual = _sum_values(
                result["rows"], result["headers"].index(check["output_field"])
            )
            checks.append(
                {
                    **check,
                    "scope": "matched_source_rows",
                    "expected": expected,
                    "actual": actual,
                    "passed": _equal_result([[expected]], [[actual]]),
                }
            )

    # These are semantic invariance probes over copies, not simulations of Excel
    # calculation and not evidence that the selected business metric is correct.
    reversed_columns = {
        **source,
        "headers": list(reversed(headers)),
        "rows": [list(reversed(row)) for row in rows],
    }
    reversed_rows = {**source, "rows": list(reversed(rows))}
    probes = []
    for name, variant in (
        ("column_order", reversed_columns),
        ("row_order", reversed_rows),
    ):
        other = _aggregate(variant, recipe)
        probes.append(
            {
                "type": name,
                "passed": result["headers"] == other["headers"]
                and _equal_result(result["rows"], other["rows"]),
            }
        )

    group_indexes = [headers.index(field) for field in recipe["group_by"]]
    members: dict[tuple, list[int]] = {}
    for row_number, row in matched:
        key = tuple(row[index] for index in group_indexes)
        members.setdefault(key, []).append(row_number)
    lineage = []
    for row in result["rows"][:sample_limit]:
        key = tuple(row[: len(group_indexes)])
        contributing = members.get(key, [])
        lineage.append(
            {
                "group": dict(zip(recipe["group_by"], key)),
                "source_row_count": len(contributing),
                "source_rows": contributing[:sample_limit],
                "truncated": len(contributing) > sample_limit,
            }
        )
    return {
        "passed": all(item["passed"] for item in checks + probes),
        "checks": checks,
        "invariance_probes": probes,
        "numeric_tolerance": {"relative": NUMERIC_REL_TOL, "absolute": NUMERIC_ABS_TOL},
        "floating_point_output": "normalized_to_openpyxl_numeric_serialization",
        "coverage": "complete bounded source and output; lineage samples only",
        "lineage_sample": lineage,
        "lineage_truncated": len(result["rows"]) > sample_limit,
    }


def _overlap(left: tuple, right: tuple) -> bool:
    a, b, c, d = left
    e, f, g, h = right
    return a <= g and e <= c and b <= h and f <= d


def _compile_output(
    wb: Any, recipe: dict, source: dict, result: dict
) -> tuple[list, list]:
    output = recipe["output"]
    bounds = changeset._range_bounds(output["range_ref"], location="output.range_ref")
    c1, r1, c2, r2 = bounds
    data = [result["headers"], *result["rows"]]
    if len(data) > r2 - r1 + 1 or len(result["headers"]) > c2 - c1 + 1:
        raise ValidationError(
            "Recipe result does not fit the reserved output range; enlarge it and preview again"
        )
    for row in data:
        for value in row:
            _scalar(value, location="Output")
            if isinstance(value, str) and (
                value.startswith("=") or value in ERROR_CODES
            ):
                raise ValidationError(
                    "Recipe output cannot contain formula-like strings or Excel error literals"
                )
    sheet_name = output["sheet_name"]
    if sheet_name not in wb.sheetnames and sheet_name.casefold() in {
        s.casefold() for s in wb.sheetnames
    }:
        raise ValidationError(
            "output.sheet_name differs only in case from an existing sheet"
        )
    if sheet_name in wb.sheetnames:
        ws = workbook.require_worksheet(
            wb, sheet_name, error_cls=ValidationError, operation="recipe output"
        )
        if ws.protection.sheet:
            raise ValidationError("Recipe output worksheet is protected")
        occupied_ranges = [table.ref for table in ws.tables.values()] + [
            str(r) for r in ws.merged_cells.ranges
        ]
        for ref in occupied_ranges:
            if _overlap(
                bounds, changeset._range_bounds(ref, location="existing range")
            ):
                raise ValidationError(
                    f"Recipe output overlaps a native table or merged range: {ref}"
                )
        if not output["allow_overwrite"]:
            for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
                if any(cell.value is not None for cell in row):
                    raise ValidationError(
                        "Output contains values; set output.allow_overwrite=true for the reserved range"
                    )
    padded = [row + [None] * (c2 - c1 + 1 - len(row)) for row in data]
    padded += [[None] * (c2 - c1 + 1) for _ in range(r2 - r1 + 1 - len(padded))]
    operations = [
        {
            "tool": "write_data_to_excel",
            "args": {
                "sheet_name": sheet_name,
                "start_cell": f"{get_column_letter(c1)}{r1}",
                "data": padded,
            },
        }
    ]
    assertions = [
        {
            "type": "range_equals",
            "sheet_name": sheet_name,
            "range_ref": output["range_ref"],
            "expected": padded,
        }
    ]
    # Protect the complete source table, splitting at the existing assertion cap.
    a, b, c, d = changeset._range_bounds(
        source["range_ref"], location="source.range_ref"
    )
    chunk_rows = changeset.MAX_ASSERTED_RANGE_CELLS // (c - a + 1)
    for start in range(b, d + 1, chunk_rows):
        end = min(start + chunk_rows - 1, d)
        assertions.append(
            {
                "type": "range_values_unchanged",
                "sheet_name": source["sheet_name"],
                "range_ref": f"{get_column_letter(a)}{start}:{get_column_letter(c)}{end}",
            }
        )
    return changeset._normalize_operations(operations), changeset._normalize_assertions(
        assertions
    )


def _assert_summary_budget(summary: dict) -> None:
    # 8k recipe metadata + the 40k ChangeSet budget leaves 2k for its receipt
    # and the MCP envelope. Count exactly the escaping used by the server.
    size = len(
        json.dumps(summary, separators=(",", ":"), ensure_ascii=True, allow_nan=False)
    )
    if size > MAX_SUMMARY_BYTES:
        raise ResponseTooLargeError(
            "Recipe summary exceeds its pre-commit response budget",
            estimated_size=size,
            limit=MAX_SUMMARY_BYTES,
            hints=["reduce sample_limit or the number of fields/checks"],
        )


def run_workbook_recipe(
    filepath: str,
    recipe: dict,
    mode: str = "preview",
    expected_workbook_sha256: str | None = None,
    recipe_token: str | None = None,
    sample_limit: int = 5,
    workbook_contract: dict | None = None,
) -> dict:
    """Compile a native-table recipe from a stable snapshot and use ChangeSet persistence."""
    if not isinstance(mode, str) or mode not in {"preview", "commit"}:
        raise ValidationError("mode must be preview or commit")
    if type(sample_limit) is not int or not 1 <= sample_limit <= 10:
        raise ValidationError("sample_limit must be an integer from 1 to 10")
    expected = changeset._normalize_sha256(
        expected_workbook_sha256, required=mode == "commit"
    )
    if recipe_token is not None and (
        not isinstance(recipe_token, str)
        or re.fullmatch(r"recipe_v1_[0-9a-f]{64}", recipe_token) is None
    ):
        raise ValidationError(
            "recipe_token must be a recipe_v1_ SHA-256 token from preview"
        )
    if mode == "commit" and recipe_token is None:
        raise ValidationError("recipe_token is required in commit mode")
    normalized = _normalize_recipe(recipe)
    requested = normalized
    contract = None
    if workbook_contract is not None:
        from .contracts import normalize_contract

        contract = normalize_contract(workbook_contract)
    path = workbook._canonical_workbook_path(filepath, must_exist=True)
    if path.suffix.lower() != ".xlsx":
        raise ValidationError("filepath must end with .xlsx")
    with _recipe_workspace() as temp:
        baseline, candidate = temp / "baseline.xlsx", temp / "candidate.xlsx"
        with workbook._exclusive_workbook_lock(path):
            digest = changeset._stage_workbook(path, baseline, candidate)
        if expected is not None and expected != digest:
            raise PreconditionFailedError(
                "Workbook changed; preview the recipe again",
                code="stale_workbook",
                suggested_next_tool="run_workbook_recipe",
            )
        with safe_workbook(str(baseline)) as wb:
            contract_report = None
            if contract is not None:
                from .contracts import compile_contract_recipe

                normalized, contract_report = compile_contract_recipe(
                    wb, normalized, contract, min(sample_limit, 5)
                )
            source = _load_table(wb, normalized)
            result = _aggregate(source, normalized)
            output_result = result
            if contract is not None:
                output_result = {
                    **result,
                    "headers": [
                        *requested["group_by"],
                        *[m["as"] for m in requested["metrics"]],
                    ],
                }
            operations, assertions = _compile_output(
                wb, normalized, source, output_result
            )
            if contract_report is not None:
                assertions = changeset._normalize_assertions(
                    [
                        *assertions,
                        {
                            "type": "table_contract_matches",
                            "contract": contract,
                            "expected_binding_sha256": contract_report[
                                "binding_sha256"
                            ],
                        },
                    ]
                )
        verification = _verify_recipe(source, normalized, result, sample_limit)
        plan = changeset._plan_payload(
            source=path,
            source_sha256=digest,
            operations=operations,
            assertions=assertions,
            create_snapshot=True,
            snapshot_filepath=None,
        )
        changeset_token = changeset._changeset_token(plan)
        token_input = {
            "engine_version": RECIPE_ENGINE_VERSION,
            "recipe": normalized,
            "changeset_token": changeset_token,
        }
        if contract is not None:
            from .contracts import CONTRACT_ENGINE_VERSION

            token_input.update(
                workbook_contract=contract,
                contract_engine_version=CONTRACT_ENGINE_VERSION,
            )
        token = (
            "recipe_v1_"
            + hashlib.sha256(
                changeset._canonical_json(token_input).encode("utf-8")
            ).hexdigest()
        )
        if recipe_token is not None and not hmac.compare_digest(recipe_token, token):
            raise PreconditionFailedError(
                "Recipe token does not match the source, recipe, or compiled output",
                code="recipe_token_mismatch",
                suggested_next_tool="run_workbook_recipe",
            )
        summary = changeset._bound_response_value(
            {
                "recipe_token": token,
                "recipe_version": 1,
                "engine_version": RECIPE_ENGINE_VERSION,
                "source": {
                    key: source[key]
                    for key in (
                        "sheet_name",
                        "table_name",
                        "range_ref",
                        "bindings",
                        "totals_rows_excluded",
                    )
                },
                "source_sha256": digest,
                "source_row_count": len(source["rows"]),
                "matched_row_count": result["matched_rows"],
                "output_row_count": len(result["rows"]),
                "output": normalized["output"],
                "headers": output_result["headers"],
                "rows_sample": result["rows"][:sample_limit],
                "truncated": len(result["rows"]) > sample_limit,
                "verification": verification,
            }
        )
        if contract_report is not None:
            summary["workbook_contract"] = changeset._bound_response_value(
                contract_report
            )
        _assert_summary_budget(summary)
        if not verification["passed"]:
            if mode == "commit":
                raise PreconditionFailedError(
                    "Recipe checks or invariance probes failed; workbook unchanged",
                    code="recipe_verification_failed",
                    details=summary,
                    suggested_next_tool="run_workbook_recipe",
                )
            return {
                "message": "Recipe verification requires attention",
                "mode": mode,
                "persisted": False,
                "ready_to_commit": False,
                "expected_workbook_sha256": digest,
                **summary,
            }

        # The compiled values never pass through the model. ChangeSet independently
        # rechecks this exact source hash under its lock before replacing the file.
        transaction = changeset.apply_workbook_changeset(
            str(path),
            operations,
            assertions,
            mode=mode,
            expected_workbook_sha256=digest,
            changeset_token=changeset_token,
            sample_limit=sample_limit,
        )
        transaction.pop("changeset_token", None)
        return {
            **transaction,
            **summary,
            "message": f"Recipe {mode} completed and verified",
        }
