"""Rebind literal native-table patches with explicit row-local read dependencies."""

from fractions import Fraction
import hashlib
import hmac
import json
import re

from openpyxl.cell.cell import ERROR_CODES, ILLEGAL_CHARACTERS_RE
from openpyxl.compat import safe_string
from openpyxl.utils import get_column_letter

from . import changeset, recipes, workbook
from .exceptions import PreconditionFailedError, ResponseTooLargeError, ValidationError
from .tables import _find_table, _table_data_bounds

REBASE_VERSION = 1
MAX_PATCH_BYTES = 128_000
MAX_TARGET_CELLS = 1000
MAX_SUMMARY_CHARS = 8000


def _identity(value):
    recipes._scalar(value, location="Referenced value")
    return (
        recipes._kind(value),
        Fraction(str(value)) if type(value) in (int, float) else value,
    )


def _key(values):
    if any(v is None or isinstance(v, str) and not v.strip() for v in values):
        raise ValidationError("Every key component must be nonblank")
    return tuple(_identity(v) for v in values)


def _output(value):
    recipes._scalar(value, location="Patch output")
    if isinstance(value, str):
        try:
            value.encode("utf-8")
        except UnicodeError as exc:
            raise ValidationError("Patch output must be valid Unicode") from exc
        if (
            not value
            or len(value) > 32767
            or ILLEGAL_CHARACTERS_RE.search(value)
            or value.startswith("=")
            or value in ERROR_CODES
        ):
            raise ValidationError(
                "Patch output must be literal Excel text (1-32767 characters), not a "
                "formula/error/control string; use null to clear a cell"
            )
    if type(value) in (int, float):
        try:
            serialized = safe_string(value)
            preserved = (
                Fraction(serialized) == Fraction(str(value))
                and float(serialized) == value
            )
        except (OverflowError, ValueError):
            preserved = False
        if not preserved:
            raise ValidationError(
                "Patch number cannot be preserved exactly by Excel serialization"
            )
    return value


def _contract(patch):
    patch = changeset._json_copy(patch, location="patch")
    try:
        size = len(changeset._canonical_json(patch).encode("utf-8"))
    except UnicodeError as exc:
        raise ValidationError("patch must contain valid Unicode") from exc
    if size > MAX_PATCH_BYTES:
        raise ValidationError(f"patch exceeds {MAX_PATCH_BYTES} bytes")
    recipes._object(
        patch,
        allowed={"version", "source", "key_fields", "read_fields", "updates"},
        required={"version", "source", "key_fields", "read_fields", "updates"},
        location="patch",
    )
    if type(patch["version"]) is not int or patch["version"] != 1:
        raise ValidationError("patch.version must be integer 1")
    source = recipes._object(
        patch["source"],
        allowed={"table_name", "sheet_name"},
        required={"table_name"},
        location="source",
    )
    for name, value in source.items():
        recipes._name(value, f"source.{name}")
    keys = recipes._field_list(patch["key_fields"], "key_fields")
    reads = recipes._field_list(patch["read_fields"], "read_fields")
    if not keys:
        raise ValidationError("key_fields must not be empty")
    updates = patch["updates"]
    if not isinstance(updates, list) or not 1 <= len(updates) <= MAX_TARGET_CELLS:
        raise ValidationError("updates must contain 1-1000 keyed row updates")
    seen, fields, count = set(), set(keys + reads), 0
    for update in updates:
        recipes._object(
            update,
            allowed={"key", "values"},
            required={"key", "values"},
            location="update",
        )
        recipes._object(
            update["key"], allowed=set(keys), required=set(keys), location="update.key"
        )
        identity = _key([update["key"][field] for field in keys])
        if identity in seen:
            raise ValidationError("Duplicate update key")
        seen.add(identity)
        values = update["values"]
        if not isinstance(values, dict) or not values:
            raise ValidationError("update.values must be a nonempty object")
        for field, value in values.items():
            recipes._name(field, "update field")
            if field in keys or field in reads:
                raise ValidationError(
                    "Write fields must be disjoint from key_fields and read_fields"
                )
            _output(value)
        fields.update(values)
        count += len(values)
    if count > MAX_TARGET_CELLS or len(fields) > 64:
        raise ValidationError("Patch exceeds 1000 target cells or 64 referenced fields")
    return patch, sorted(fields)


def _load(wb, patch, fields):
    sheet, ws, table = _find_table(wb, **patch["source"])
    c1, r1, c2, r2 = changeset._range_bounds(table.ref, location="table range")
    if table.headerRowCount != 1 or int(table.totalsRowCount or 0) not in (0, 1):
        raise ValidationError(
            "Patch requires one header row and at most one totals row"
        )
    if r2 - r1 < int(table.totalsRowCount or 0):
        raise ValidationError("Inconsistent table totals-row bounds")
    if (c2 - c1 + 1) * (r2 - r1 + 1) > 50_000 or c2 - c1 + 1 > 64:
        raise ValidationError("Patch table exceeds 50000 cells or 64 columns")
    headers = [ws.cell(r1, col).value for col in range(c1, c2 + 1)]
    for header in headers:
        recipes._name(header, "table header")
    if len({h.strip().casefold() for h in headers}) != len(headers):
        raise ValidationError("Table has ambiguous duplicate headers")
    if set(fields) - set(headers):
        raise ValidationError("Referenced fields missing; exact header names required")
    columns = {field: c1 + headers.index(field) for field in fields}
    if (
        len(table.tableColumns) != len(headers)
        or [c.name for c in table.tableColumns] != headers
    ):
        raise ValidationError("Table column metadata does not match worksheet headers")
    if any(
        table.tableColumns[headers.index(f)].calculatedColumnFormula is not None
        for f in fields
    ):
        raise ValidationError("Referenced fields cannot be calculated table columns")
    _, _, _, _, start, end = _table_data_bounds(table)
    records = {}
    for row in range(start, end + 1):
        values = []
        for field in patch["key_fields"]:
            cell = ws.cell(row, columns[field])
            if cell.data_type in {"f", "e"}:
                raise ValidationError(
                    "Keys must be literal values, not formulas/errors"
                )
            values.append(cell.value)
        key = _key(values)
        if key in records:
            raise ValidationError("Table contains duplicate keys")
        records[key] = row
    return ws, headers, columns, records, table


def _literal(ws, row, col):
    cell = ws.cell(row, col)
    if cell.data_type in {"f", "e"}:
        raise ValidationError(
            "Referenced patch cells must be literals, not formulas/errors"
        )
    return recipes._scalar(cell.value, location="Referenced patch cell")


def _compile(base_wb, current_wb, patch, fields, sample_limit):
    before = _load(base_wb, patch, fields)
    after = _load(current_wb, patch, fields)
    old_ws, old_headers, old_cols, old_rows, _ = before
    ws, headers, cols, rows, table = after
    if set(old_headers) != set(headers):
        raise ValidationError(
            "Table header set changed; only column reordering is supported"
        )
    if ws.protection.sheet:
        raise ValidationError("Patch target worksheet is protected")
    counts = {
        name: 0
        for name in (
            "apply",
            "already_satisfied",
            "write_conflict",
            "stale_precondition",
            "missing_row",
        )
    }
    examples, writes = [], []
    for update in patch["updates"]:
        identity = _key([update["key"][field] for field in patch["key_fields"]])
        if identity not in old_rows:
            raise ValidationError("Update key does not exist in the pinned base table")
        old_row, row = old_rows[identity], rows.get(identity)
        stale = (
            []
            if row is None
            else [
                field
                for field in patch["read_fields"]
                if _identity(_literal(old_ws, old_row, old_cols[field]))
                != _identity(_literal(ws, row, cols[field]))
            ]
        )
        for field, desired in update["values"].items():
            old = _literal(old_ws, old_row, old_cols[field])
            current = None if row is None else _literal(ws, row, cols[field])
            if row is not None:
                for merged in ws.merged_cells.ranges:
                    if ws.cell(row, cols[field]).coordinate in merged:
                        raise ValidationError("Patch target overlaps merged cells")
            if row is None:
                state = "missing_row"
            elif stale:
                state = "stale_precondition"
            elif _identity(current) == _identity(desired):
                state = "already_satisfied"
            elif _identity(current) == _identity(old):
                state = "apply"
                writes.append((cols[field], row, desired))
            else:
                state = "write_conflict"
            counts[state] += 1
            examples.append(
                {
                    "key": update["key"],
                    "field": field,
                    "status": state,
                    "base": old,
                    "current": current,
                    "desired": desired,
                    "changed_read_fields": stale,
                    "current_cell": None
                    if row is None
                    else ws.cell(row, cols[field]).coordinate,
                }
            )
    blocked = any(
        counts[name] for name in ("write_conflict", "stale_precondition", "missing_row")
    )
    operations, assertions = [], []
    if not blocked:
        spans = []
        for col, row, value in sorted(writes):
            if (
                spans
                and spans[-1][0] == col
                and spans[-1][1] + len(spans[-1][2]) == row
            ):
                spans[-1][2].append([value])
            else:
                spans.append((col, row, [[value]]))
        if len(spans) > 50:
            raise ValidationError(
                "Patch exceeds 50 contiguous column spans; split into smaller patches"
            )
        for col, row, data in spans:
            cell = f"{get_column_letter(col)}{row}"
            operations.append(
                {
                    "tool": "write_data_to_excel",
                    "args": {"sheet_name": ws.title, "start_cell": cell, "data": data},
                }
            )
            assertions.append(
                {
                    "type": "range_equals",
                    "sheet_name": ws.title,
                    "range_ref": f"{cell}:{get_column_letter(col)}{row + len(data) - 1}",
                    "expected": data,
                }
            )
    priority = {
        "write_conflict": 0,
        "stale_precondition": 0,
        "missing_row": 0,
        "apply": 1,
        "already_satisfied": 2,
    }
    examples.sort(key=lambda item: priority[item["status"]])
    examples = examples[:sample_limit]
    return (
        operations,
        assertions,
        {
            "counts": counts,
            "sample": examples,
            "truncated": sum(counts.values()) > len(examples),
            "blocked": blocked,
            "read_fields": patch["read_fields"],
            "bindings": {
                field: f"{get_column_letter(col)}{changeset._range_bounds(table.ref, location='table range')[1]}"
                for field, col in cols.items()
            },
        },
    )


def _summary_budget(summary):
    size = len(
        json.dumps(summary, ensure_ascii=True, allow_nan=False, separators=(",", ":"))
    )
    if size > MAX_SUMMARY_CHARS:
        raise ResponseTooLargeError(
            "Rebase summary exceeds response budget",
            estimated_size=size,
            limit=MAX_SUMMARY_CHARS,
            hints=["reduce sample_limit or patch size"],
        )


def apply_table_patch(
    base_filepath: str,
    filepath: str,
    patch: dict,
    expected_base_sha256: str,
    mode: str = "preview",
    expected_workbook_sha256: str | None = None,
    rebase_token: str | None = None,
    sample_limit: int = 5,
) -> dict:
    """Preview a three-way keyed patch; commit only the exact fresh preview."""
    if not isinstance(mode, str) or mode not in {"preview", "commit"}:
        raise ValidationError("mode must be preview or commit")
    if type(sample_limit) is not int or not 1 <= sample_limit <= 10:
        raise ValidationError("sample_limit must be an integer from 1 to 10")
    if expected_base_sha256 is None:
        raise ValidationError("expected_base_sha256 is required on every call")
    base_sha = changeset._normalize_sha256(expected_base_sha256, required=True)
    expected = changeset._normalize_sha256(
        expected_workbook_sha256, required=mode == "commit"
    )
    if rebase_token is not None and (
        not isinstance(rebase_token, str)
        or re.fullmatch(r"rebase_v1_[0-9a-f]{64}", rebase_token) is None
    ):
        raise ValidationError("rebase_token must be a rebase_v1_ token from preview")
    if mode == "commit" and rebase_token is None:
        raise ValidationError("rebase_token is required in commit mode")
    patch, fields = _contract(patch)
    paths = [
        workbook._canonical_workbook_path(p, must_exist=True)
        for p in (base_filepath, filepath)
    ]
    if any(p.suffix.lower() != ".xlsx" or not p.is_file() for p in paths):
        raise ValidationError("Base and current must be regular .xlsx files")
    if paths[0].samefile(paths[1]):
        raise ValidationError("Base and current must be distinct files, not aliases")
    with recipes._recipe_workspace() as temp:
        captures, digests = [], []
        for i, path in enumerate(paths):
            frozen = temp / f"capture-{i}.xlsx"
            with workbook._exclusive_workbook_lock(path):
                digest = changeset._stage_workbook(
                    path, frozen, temp / f"unused-{i}.xlsx"
                )
            pin = base_sha if i == 0 else expected
            if pin is not None and pin != digest:
                raise PreconditionFailedError(
                    "Pinned workbook changed; create a fresh preview",
                    code="stale_workbook",
                    details={
                        "source": "base" if i == 0 else "current",
                        "expected_sha256": pin,
                        "actual_sha256": digest,
                    },
                    suggested_next_tool="apply_table_patch",
                )
            captures.append(frozen)
            digests.append(digest)
        with workbook.safe_workbook(str(captures[0])) as base_wb:
            with workbook.safe_workbook(str(captures[1])) as current_wb:
                operations, assertions, summary = _compile(
                    base_wb, current_wb, patch, fields, sample_limit
                )
    if operations:
        operations = changeset._normalize_operations(operations)
        assertions = changeset._normalize_assertions(assertions)
    plan = changeset._plan_payload(
        source=paths[1],
        source_sha256=digests[1],
        operations=operations,
        assertions=assertions,
        create_snapshot=True,
        snapshot_filepath=None,
    )
    cs_token = changeset._changeset_token(plan)
    identity = {
        "version": REBASE_VERSION,
        "base_path": str(paths[0]),
        "base_sha256": base_sha,
        "patch": patch,
        "changeset_token": cs_token,
    }
    token = (
        "rebase_v1_"
        + hashlib.sha256(
            changeset._canonical_json(identity).encode("utf-8")
        ).hexdigest()
    )
    if rebase_token is not None and not hmac.compare_digest(rebase_token, token):
        raise PreconditionFailedError(
            "Rebase token does not match the base/current/patch plan",
            code="rebase_token_mismatch",
            suggested_next_tool="apply_table_patch",
        )
    summary = changeset._bound_response_value(
        {
            "rebase_version": REBASE_VERSION,
            "expected_base_sha256": base_sha,
            "expected_workbook_sha256": digests[1],
            "rebase_token": token,
            **summary,
            "verification": {
                "coverage": "complete_keys_and_declared_row_local_dependencies",
                "undeclared_dependencies_checked": False,
                "cross_file_atomic": False,
            },
        }
    )
    _summary_budget(summary)
    if summary["blocked"]:
        if mode == "commit":
            raise PreconditionFailedError(
                "Patch has conflicts; no changes written",
                code="rebase_conflict",
                details=summary,
                suggested_next_tool="apply_table_patch",
            )
        return {**summary, "mode": mode, "persisted": False, "ready_to_commit": False}
    if not operations:
        return {
            **summary,
            "mode": mode,
            "persisted": False,
            "ready_to_commit": True,
            "message": "All target values already satisfied in the captured current version",
        }
    transaction = changeset.apply_workbook_changeset(
        str(paths[1]),
        operations,
        assertions,
        mode=mode,
        expected_workbook_sha256=digests[1],
        changeset_token=cs_token,
        sample_limit=sample_limit,
    )
    transaction.pop("changeset_token", None)
    return {
        **transaction,
        **summary,
        "message": f"Table patch {mode} completed and verified",
    }
