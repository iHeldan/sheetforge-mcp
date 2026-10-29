"""Actual stdio contract handoff, rejection and complete report readback."""

import asyncio
import copy
from decimal import Decimal
import json
from pathlib import Path
import sys

from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table


def exercise_contract_workflow(tmp_path, *, installed=False):
    root = Path(__file__).resolve().parents[1]
    contract = json.loads((root / "docs/examples/net-sales.contract.json").read_text())
    path = tmp_path / "monthly.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Status", "Currency", "NetSales", "OrderId"])
    expected = Decimal(0)
    for i in range(120):
        status = "cancelled" if i % 3 == 0 else "paid"
        amount = (i + 1) / 10
        ws.append([status, "EUR", amount, i])
        if status == "paid":
            expected += Decimal(str(amount))
    ws.add_table(Table(displayName="Sales", ref="A1:D121"))
    ws["G1"] = "=1+1"
    wb.save(path)
    wb.close()
    recipe = {
        "version": 1,
        "source": {"table_name": "Sales"},
        "group_by": ["currency"],
        "metrics": [{"op": "sum", "field": "net", "as": "Revenue"}],
        "output": {"sheet_name": "Report", "range_ref": "A1:B5"},
    }
    parameters = StdioServerParameters(
        command=sys.executable,
        args=["-m", "excel_mcp", "stdio"],
        cwd=tmp_path if installed else root,
        env={"PYTHONPATH": "" if installed else str(root / "src")},
    )
    sizes = []

    async def exercise():
        with (tmp_path / "stderr.log").open("w") as log:
            async with stdio_client(parameters, errlog=log) as (read, write):
                async with ClientSession(read, write) as session:
                    await session.initialize()
                    inventory = (await session.list_tools()).tools
                    assert len(inventory) == 83
                    tool = next(
                        t for t in inventory if t.name == "inspect_workbook_contract"
                    )
                    assert tool.outputSchema is None and tool.annotations.readOnlyHint
                    assert set(tool.inputSchema["required"]) == {"filepath", "contract"}
                    runner = next(
                        t for t in inventory if t.name == "run_workbook_recipe"
                    )
                    assert "workbook_contract" in runner.inputSchema["properties"]

                    async def call(name, args):
                        response = await session.call_tool(
                            name, {"filepath": str(path), **args}
                        )
                        wire = "".join(
                            c.text for c in response.content if c.type == "text"
                        )
                        assert len(wire) <= 50000
                        sizes.append(len(wire))
                        return json.loads(wire)

                    before, mtime = path.read_bytes(), path.stat().st_mtime_ns
                    inspection = await call(
                        "inspect_workbook_contract", {"contract": contract}
                    )
                    assert inspection["data"]["passed"]
                    assert inspection["data"]["source_row_count"] == 120
                    assert (
                        path.read_bytes() == before and path.stat().st_mtime_ns == mtime
                    )
                    invalid = await call(
                        "inspect_workbook_contract",
                        {"contract": {**contract, "force": True}},
                    )
                    assert not invalid["ok"] and path.read_bytes() == before
                    args = {
                        "recipe": recipe,
                        "workbook_contract": contract,
                        "sample_limit": 1,
                    }
                    preview = (await call("run_workbook_recipe", args))["data"]
                    assert preview["ready_to_commit"] and path.read_bytes() == before
                    pin = {
                        "mode": "commit",
                        "recipe_token": preview["recipe_token"],
                        "expected_workbook_sha256": preview["expected_workbook_sha256"],
                    }
                    modified = copy.deepcopy(contract)
                    modified["row_grain"] = "Changed declaration"
                    rejected = await call(
                        "run_workbook_recipe",
                        {**args, **pin, "workbook_contract": modified},
                    )
                    assert not rejected["ok"] and path.read_bytes() == before
                    result = await call("run_workbook_recipe", {**args, **pin})
                    assert result["ok"] and result["data"]["persisted"]
                    wb = load_workbook(path)
                    assert list(wb["Report"].values) == [
                        ("currency", "Revenue"),
                        ("EUR", float(expected)),
                    ]
                    assert wb["Data"]["G1"].value == "=1+1"
                    assert [
                        list(row)
                        for row in wb["Data"].iter_rows(
                            min_row=1, max_row=121, max_col=4, values_only=True
                        )
                    ] == [
                        ["Status", "Currency", "NetSales", "OrderId"],
                        *[
                            [
                                "cancelled" if i % 3 == 0 else "paid",
                                "EUR",
                                (i + 1) / 10,
                                i,
                            ]
                            for i in range(120)
                        ],
                    ]
                    wb.close()
                    assert not (await call("run_workbook_recipe", {**args, **pin}))[
                        "ok"
                    ]
                    wb = load_workbook(path)
                    wb["Data"]["B2"] = "USD"
                    wb.save(path)
                    wb.close()
                    changed = path.read_bytes()
                    assert not (
                        await call("inspect_workbook_contract", {"contract": contract})
                    )["data"]["passed"]
                    assert not (
                        await call(
                            "run_workbook_recipe",
                            {
                                **args,
                                "recipe": {
                                    **recipe,
                                    "output": {
                                        **recipe["output"],
                                        "allow_overwrite": True,
                                    },
                                },
                            },
                        )
                    )["ok"]
                    assert path.read_bytes() == changed
                    assert sum(sizes) <= 100000

    asyncio.run(asyncio.wait_for(exercise(), timeout=45))
    return {
        "tools": 83,
        "calls": len(sizes),
        "response_chars": sum(sizes),
        "installed": installed,
    }


def test_stdio_contract_handoff_and_guarded_recipe(tmp_path):
    exercise_contract_workflow(tmp_path)
