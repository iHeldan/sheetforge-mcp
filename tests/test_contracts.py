import copy
import json

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from excel_mcp import changeset, contracts, recipes, server
from excel_mcp.exceptions import (
    DataError,
    PreconditionFailedError,
    ResponseTooLargeError,
    ValidationError,
    WorkbookError,
)


@pytest.fixture
def contract():
    return {
        "version": 1,
        "name": "net-sales",
        "source": {"table_name": "Sales"},
        "row_grain": "One row per order",
        "key_fields": ["id"],
        "fields": [
            {"name": "id", "headers": ["Id", "OrderId"], "type": "integer"},
            {
                "name": "net",
                "headers": ["Net", "NetSales"],
                "type": "number",
                "min": 0,
                "meaning": "Net amount excluding tax",
            },
            {
                "name": "currency",
                "headers": ["Currency"],
                "type": "string",
                "allowed_values": ["EUR"],
            },
            {
                "name": "status",
                "headers": ["Status"],
                "type": "string",
                "allowed_values": ["paid", "cancelled"],
            },
        ],
        "filters": [
            {"field": "status", "op": "eq", "value": "paid", "case_sensitive": True}
        ],
    }


@pytest.fixture
def book(tmp_path):
    path = tmp_path / "sales.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in [
        ["Id", "Net", "Currency", "Status"],
        [1, 10, "EUR", "paid"],
        [2, 20, "EUR", "cancelled"],
        [3, 30, "EUR", "paid"],
    ]:
        ws.append(row)
    ws.add_table(Table(displayName="Sales", ref="A1:D4"))
    ws["G1"] = "=1+1"
    wb.save(path)
    wb.close()
    return path


@pytest.fixture
def recipe():
    return {
        "version": 1,
        "source": {"table_name": "Sales"},
        "group_by": ["currency"],
        "metrics": [{"op": "sum", "field": "net", "as": "Revenue"}],
        "output": {"sheet_name": "Report", "range_ref": "A1:B5"},
    }


def edit(path, action):
    wb = load_workbook(path)
    try:
        action(wb["Data"])
        wb.save(path)
    finally:
        wb.close()


def commit(path, recipe, contract, preview):
    return recipes.run_workbook_recipe(
        str(path),
        recipe,
        mode="commit",
        workbook_contract=contract,
        recipe_token=preview["recipe_token"],
        expected_workbook_sha256=preview["expected_workbook_sha256"],
    )


def test_inspect_read_only_full_coverage_and_declarations(book, contract):
    before, mtime, original = (
        book.read_bytes(),
        book.stat().st_mtime_ns,
        copy.deepcopy(contract),
    )
    result = contracts.inspect_workbook_contract(str(book), contract, sample_limit=1)
    assert result["passed"] and result["status"] == "valid"
    assert result["source_row_count"] == 3
    assert result["bindings"]["net"] == "Net"
    assert result["declared"]["meanings"]["net"] == "Net amount excluding tax"
    assert all(c["passed"] for c in result["checked"])
    assert contract == original
    assert book.read_bytes() == before and book.stat().st_mtime_ns == mtime
    assert not list(book.parent.glob("*.sheetforge-before-*"))


@pytest.mark.parametrize("header", ["OrderId", "NetSales"])
def test_approved_aliases(book, contract, header):
    column = 1 if header == "OrderId" else 2
    edit(book, lambda ws: setattr(ws.cell(1, column), "value", header))
    assert contracts.inspect_workbook_contract(str(book), contract)["passed"]


@pytest.mark.parametrize(
    "header,status", [("Unknown", "missing"), ("NetSales", "ambiguous")]
)
def test_no_guessing_missing_or_ambiguous_headers(
    book, contract, recipe, header, status
):
    if status == "missing":
        edit(book, lambda ws: setattr(ws["B1"], "value", header))
    else:

        def ambiguous(ws):
            ws["E1"] = header
            ws.tables["Sales"].ref = "A1:E4"

        edit(book, ambiguous)
    before = book.read_bytes()
    report = contracts.inspect_workbook_contract(str(book), contract)
    assert report["status"] == "unresolved" and not report["passed"]
    assert report["checked"] == [] and "no value rules evaluated" in report["coverage"]
    assert report["unresolved"][0]["reason"] == status
    with pytest.raises(PreconditionFailedError, match="unresolved"):
        recipes.run_workbook_recipe(str(book), recipe, workbook_contract=contract)
    assert book.read_bytes() == before


@pytest.mark.parametrize(
    "cell,value",
    [
        ("C3", "USD"),
        ("A3", 1),
        ("A3", None),
        ("B3", -1),
        ("D3", "unknown"),
        ("B3", None),
    ],
)
def test_invalid_rows_cannot_be_hidden_by_mandatory_filter(
    book, contract, recipe, cell, value
):
    edit(book, lambda ws: setattr(ws[cell], "value", value))
    before = book.read_bytes()
    report = contracts.inspect_workbook_contract(str(book), contract, sample_limit=1)
    assert not report["passed"] and report["status"] == "violated"
    assert any(c.get("row_examples") == [3] for c in report["checked"])
    with pytest.raises(PreconditionFailedError):
        recipes.run_workbook_recipe(str(book), recipe, workbook_contract=contract)
    assert book.read_bytes() == before


def test_contract_recipe_commit_logical_headers_and_required_filters(
    book, contract, recipe
):
    original = copy.deepcopy(recipe)
    preview = recipes.run_workbook_recipe(str(book), recipe, workbook_contract=contract)
    assert preview["ready_to_commit"]
    assert preview["headers"] == ["currency", "Revenue"]
    assert preview["rows_sample"] == [["EUR", 40]]
    result = commit(book, recipe, contract, preview)
    assert result["persisted"] and recipe == original
    wb = load_workbook(book)
    assert list(wb["Report"].values)[:2] == [("currency", "Revenue"), ("EUR", 40)]
    assert wb["Data"]["G1"].value == "=1+1"
    wb.close()
    with pytest.raises(PreconditionFailedError, match="Workbook changed"):
        commit(book, recipe, contract, preview)


@pytest.mark.parametrize("change", ["meaning", "filter", "min", "remove"])
def test_contract_is_bound_to_recipe_token(book, contract, recipe, change):
    preview = recipes.run_workbook_recipe(str(book), recipe, workbook_contract=contract)
    modified = copy.deepcopy(contract)
    if change == "meaning":
        modified["row_grain"] = "A different declaration"
    elif change == "filter":
        modified["filters"] = []
    elif change == "min":
        modified["fields"][1]["min"] = -1
    else:
        modified = None
        recipe = {
            **recipe,
            "group_by": ["Currency"],
            "metrics": [{"op": "sum", "field": "Net", "as": "Revenue"}],
        }
    before = book.read_bytes()
    with pytest.raises(PreconditionFailedError, match="token"):
        commit(book, recipe, modified, preview)
    assert book.read_bytes() == before


@pytest.mark.parametrize(
    "update",
    [
        {"version": True},
        {"version": 2},
        {"code": "eval"},
        {"fields": []},
        {"min_rows": True},
        {"key_fields": []},
        {"key_fields": ["missing"]},
        {"filters": [{"field": "net", "op": "not-an-op", "value": 1}]},
        {"filters": [{"field": "net", "op": "eq"}]},
        {"filters": [{"field": "net", "op": "in"}]},
        {"source": {"table_name": "Sales", "url": "https://example.invalid"}},
    ],
)
def test_closed_contract_grammar(contract, update):
    contract.update(update)
    with pytest.raises((ValidationError, DataError)):
        contracts.normalize_contract(contract)


@pytest.mark.parametrize(
    "update",
    [
        {"headers": []},
        {"headers": ["X"] * 9},
        {"headers": ["Currency"]},
        {"nullable": 1},
        {"type": "datetime"},
        {"type": []},
        {"allowed_values": []},
        {"allowed_values": [None]},
        {"allowed_values": [True]},
        {"min": True},
        {"min": 10, "max": 5},
        {"min": float("inf")},
    ],
)
def test_closed_field_grammar(contract, update):
    contract["fields"][1].update(update)
    with pytest.raises(ValidationError):
        contracts.normalize_contract(contract)


@pytest.mark.parametrize("sample", [0, 6, True, "1"])
def test_sample_limits(book, contract, sample):
    with pytest.raises(ValidationError):
        contracts.inspect_workbook_contract(str(book), contract, sample)


@pytest.mark.parametrize("value", ["=1+1", "#REF!", "not numeric", True])
def test_no_formula_or_type_coercion(book, contract, value):
    edit(book, lambda ws: setattr(ws["B2"], "value", value))
    before = book.read_bytes()
    try:
        result = contracts.inspect_workbook_contract(str(book), contract)
        assert not result["passed"]
    except ValidationError:
        pass
    assert book.read_bytes() == before


def test_nullable_and_bounds(book, contract):
    contract["fields"][1].update(nullable=True, max=30)
    edit(book, lambda ws: setattr(ws["B2"], "value", None))
    assert contracts.inspect_workbook_contract(str(book), contract)["passed"]
    contract["fields"][1]["max"] = 29
    assert not contracts.inspect_workbook_contract(str(book), contract)["passed"]


def test_no_physical_field_or_different_source_bypass(book, contract, recipe):
    recipe["metrics"][0]["field"] = "Net"
    with pytest.raises(ValidationError, match="logical"):
        recipes.run_workbook_recipe(str(book), recipe, workbook_contract=contract)
    recipe["source"]["sheet_name"] = "Data"
    with pytest.raises(ValidationError, match="selectors"):
        recipes.run_workbook_recipe(str(book), recipe, workbook_contract=contract)


def test_composite_key_and_whitespace_blank(book, contract):
    contract["key_fields"] = ["currency", "id"]
    assert contracts.inspect_workbook_contract(str(book), contract)["passed"]
    contract["fields"][0]["type"] = "string"

    def strings(ws):
        for r in range(2, 5):
            ws.cell(r, 1).value = str(r)
        ws["A4"] = "  "

    edit(book, strings)
    assert not contracts.inspect_workbook_contract(str(book), contract)["passed"]


def test_chart_sheet_and_transport_containment(book, contract, monkeypatch, tmp_path):
    wb = load_workbook(book)
    wb.create_chartsheet("Charts")
    # An empty chartsheet is not an openpyxl-roundtrippable fixture; the explicit
    # sheet guard is exercised in memory before any cell-grid access.
    normalized = contracts.normalize_contract(contract)
    normalized["source"]["sheet_name"] = "Charts"
    with pytest.raises((DataError, ValidationError)):
        contracts.inspect_table_contract(wb, normalized, 1)
    wb.close()
    monkeypatch.setattr(server, "EXCEL_FILES_PATH", str(tmp_path / "contained"))
    result = json.loads(server.inspect_workbook_contract(str(book), contract))
    assert not result["ok"]


def test_contract_assertion_pre_and_post_replace_rollback(
    book, contract, recipe, monkeypatch
):
    preview = recipes.run_workbook_recipe(str(book), recipe, workbook_contract=contract)
    original = book.read_bytes()
    evaluate = changeset._evaluate_assertions
    calls = []

    def fail_saved(baseline, candidate, assertions, **kwargs):
        calls.append(candidate)
        if candidate == book:
            edit(book, lambda ws: setattr(ws["C2"], "value", "USD"))
        return evaluate(baseline, candidate, assertions, **kwargs)

    monkeypatch.setattr(changeset, "_evaluate_assertions", fail_saved)
    with pytest.raises((PreconditionFailedError, WorkbookError)):
        commit(book, recipe, contract, preview)
    assert book in calls and book.read_bytes() == original


def test_escaped_wire_budget_checked_before_mutation(
    book, contract, recipe, monkeypatch
):
    before = book.read_bytes()
    monkeypatch.setattr(recipes, "MAX_SUMMARY_BYTES", 10)
    with pytest.raises(ResponseTooLargeError):
        recipes.run_workbook_recipe(str(book), recipe, workbook_contract=contract)
    assert book.read_bytes() == before
    monkeypatch.setattr(contracts, "MAX_CONTRACT_SUMMARY", 10)
    with pytest.raises(ResponseTooLargeError):
        contracts.inspect_workbook_contract(str(book), contract)


def test_unicode_and_size_validation(contract):
    contract["row_grain"] = "\ud800"
    with pytest.raises(ValidationError, match="Unicode"):
        contracts.normalize_contract(contract)
    with pytest.raises(ValidationError, match="32000"):
        contracts.normalize_contract({"large": "x" * 32000})


def test_source_limit_and_late_duplicate(book, contract):
    def rows(ws):
        for i in range(4, 121):
            ws.append([i, 1, "EUR", "paid"])
        ws.tables["Sales"].ref = "A1:D121"
        ws["A121"] = 1

    edit(book, rows)
    result = contracts.inspect_workbook_contract(str(book), contract, sample_limit=1)
    assert result["source_row_count"] == 120 and not result["passed"]
    assert result["checked"][-1]["row_examples"] == [121]
    edit(book, lambda ws: setattr(ws.tables["Sales"], "ref", "A1:D12501"))
    with pytest.raises(ValidationError, match="exceeds"):
        contracts.inspect_workbook_contract(str(book), contract)


def test_reordered_columns_offset_totals_and_row_order(tmp_path, contract, recipe):
    path = tmp_path / "offset.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for r, row in enumerate(
        [
            ["Status", "Currency", "NetSales", "OrderId"],
            ["paid", "EUR", 30, 3],
            ["cancelled", "EUR", 20, 2],
            ["paid", "EUR", 10, 1],
            ["Totals", None, "=SUM(E4:E6)", None],
        ],
        3,
    ):
        for c, value in enumerate(row, 3):
            ws.cell(r, c, value)
    ws.add_table(Table(displayName="Sales", ref="C3:F7", totalsRowCount=1))
    wb.save(path)
    wb.close()
    preview = recipes.run_workbook_recipe(str(path), recipe, workbook_contract=contract)
    assert preview["rows_sample"] == [["EUR", 40]]
    assert preview["workbook_contract"]["source_row_count"] == 3
    assert commit(path, recipe, contract, preview)["persisted"]


def test_contract_assertion_rejects_binding_drift(book, contract):
    report = contracts.inspect_workbook_contract(str(book), contract)
    assertion = changeset._normalize_assertions(
        [
            {
                "type": "table_contract_matches",
                "contract": contract,
                "expected_binding_sha256": report["binding_sha256"],
            }
        ]
    )[0]
    edit(book, lambda ws: setattr(ws["B1"], "value", "NetSales"))
    wb = load_workbook(book)
    try:
        passed, details = contracts.assert_contract(wb, assertion, 1)
        assert (
            not passed
            and details["status"] == "valid"
            and not details["binding_matches"]
        )
    finally:
        wb.close()


def test_declared_filters_and_recipe_filters_are_both_mandatory(book, contract, recipe):
    recipe["filters"] = [{"field": "net", "op": "gte", "value": 20}]
    assert recipes.run_workbook_recipe(str(book), recipe, workbook_contract=contract)[
        "rows_sample"
    ] == [["EUR", 30]]
    recipe["checks"] = [
        {"type": "unique_key", "fields": ["id"]},
        {"type": "sum_matches", "source_field": "net", "output_field": "Revenue"},
    ]
    assert recipes.run_workbook_recipe(str(book), recipe, workbook_contract=contract)[
        "verification"
    ]["passed"]


def test_empty_requires_explicit_opt_in(book, contract):
    edit(book, lambda ws: setattr(ws.tables["Sales"], "ref", "A1:D1"))
    assert not contracts.inspect_workbook_contract(str(book), contract)["passed"]
    contract["min_rows"] = 0
    result = contracts.inspect_workbook_contract(str(book), contract)
    assert result["passed"] and result["source_row_count"] == 0


def test_non_regular_input_and_symlink(book, contract, tmp_path):
    import os

    pipe = tmp_path / "pipe.xlsx"
    os.mkfifo(pipe)
    with pytest.raises((WorkbookError, ValidationError)):
        contracts.inspect_workbook_contract(str(pipe), contract)
    link = tmp_path / "alias.xlsx"
    link.symlink_to(book)
    assert contracts.inspect_workbook_contract(str(link), contract)["passed"]


def test_docs_example_parses():
    from pathlib import Path

    example = json.loads(
        (
            Path(__file__).resolve().parents[1]
            / "docs/examples/net-sales.contract.json"
        ).read_text()
    )
    assert contracts.normalize_contract(example)["version"] == 1


def test_no_duplicate_logical_output_headers(book, contract, recipe):
    recipe["metrics"][0]["as"] = "currency"
    with pytest.raises(ValidationError, match="unambiguous logical"):
        recipes.run_workbook_recipe(str(book), recipe, workbook_contract=contract)


def test_sum_check_can_reference_logical_output_group(book, contract, recipe):
    recipe["group_by"] = ["net"]
    recipe["metrics"] = [{"op": "count", "as": "Orders"}]
    recipe["checks"] = [
        {"type": "sum_matches", "source_field": "net", "output_field": "net"}
    ]
    assert recipes.run_workbook_recipe(str(book), recipe, workbook_contract=contract)[
        "verification"
    ]["passed"]


def test_capture_rejects_fifo_substituted_after_stat(book, tmp_path, monkeypatch):
    import os

    original_open = os.open
    retained = tmp_path / "retained.xlsx"

    def swap(path, flags, *args, **kwargs):
        if path == book:
            book.rename(retained)
            os.mkfifo(book)
        return original_open(path, flags, *args, **kwargs)

    monkeypatch.setattr(changeset.os, "open", swap)
    with pytest.raises(ValidationError, match="regular file"):
        changeset._stage_workbook(
            book, tmp_path / "baseline.xlsx", tmp_path / "candidate.xlsx"
        )
    assert retained.is_file()


def test_ascii_escaping_counts_toward_report_budget(monkeypatch):
    monkeypatch.setattr(contracts, "MAX_CONTRACT_SUMMARY", 500)
    with pytest.raises(ResponseTooLargeError):
        contracts._bounded_report({"declaration": "\u00e4" * 128})


def test_number_and_boolean_types_are_not_interchangeable(book, contract):
    contract["fields"][3].update(type="boolean", allowed_values=[True])
    contract["filters"] = []

    def bools(ws):
        for row in range(2, 5):
            ws.cell(row, 4).value = True

    edit(book, bools)
    assert contracts.inspect_workbook_contract(str(book), contract)["passed"]

    def numbers(ws):
        for row in range(2, 5):
            ws.cell(row, 4).value = 1

    edit(book, numbers)
    assert not contracts.inspect_workbook_contract(str(book), contract)["passed"]


def test_filter_and_field_limits(contract):
    contract["filters"] *= 21
    with pytest.raises(ValidationError, match="20 predicates"):
        contracts.normalize_contract(contract)
    contract["filters"] = []
    contract["fields"] *= 9
    with pytest.raises(ValidationError, match="32 fields"):
        contracts.normalize_contract(contract)
