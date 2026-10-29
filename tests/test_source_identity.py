"""Namespace recovery must verify identities without selecting a different source."""

import asyncio
import json
from pathlib import Path
import sys

import pytest
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table

from excel_mcp import data, server, tables, workbook
from excel_mcp.exceptions import DataError, WorkbookError


@pytest.fixture
def source(tmp_path):
    path = tmp_path / "source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in [["Id", "Amount"], [1, 10], [2, 20]]:
        ws.append(row)
    ws.move_range("A1:B3", rows=2)
    ws.add_table(Table(displayName="Sales", ref="A3:B5"))
    ws["F1"] = "Do not disclose this cell in diagnostics"
    wb.create_sheet("Report")
    chart = BarChart()
    chart.add_data(
        Reference(ws, min_col=2, min_row=3, max_row=5), titles_from_data=True
    )
    wb.create_chartsheet("Charts").add_chart(chart)
    wb.save(path)
    wb.close()
    return path


@pytest.mark.parametrize(
    "reader",
    [
        lambda p: data.read_as_table(p, "Sales"),
        lambda p: data.quick_read(p, "Sales"),
        lambda p: data.read_excel_range(p, "Sales"),
        lambda p: data.read_excel_range_with_metadata(p, "Sales"),
        lambda p: data.describe_dataset(p, sheet_name="Sales"),
    ],
)
def test_readers_preserve_type_and_explain_exact_table_match(source, reader):
    before, mtime = source.read_bytes(), source.stat().st_mtime_ns
    with pytest.raises(DataError) as caught:
        reader(str(source))
    error = caught.value
    assert type(error) is DataError and error.code == "sheet_name_is_table"
    assert error.suggested_next_tool == "list_tables"
    assert error.details["suggested_arguments"] == {"sheet_name": None}
    assert error.details["table_matches"] == [
        {"table_name": "Sales", "sheet_name": "Data", "range_ref": "A3:B5"}
    ]
    assert error.details["table_match_count"] == 1
    assert not error.details["automatic_selection"]
    assert "Do not disclose" not in json.dumps(error.details)
    assert source.read_bytes() == before and source.stat().st_mtime_ns == mtime


def test_metadata_does_not_claim_table_inventory(source):
    info = workbook.get_workbook_info(str(source), include_ranges=True)
    assert info["sheets"] == ["Data", "Report", "Charts"]
    assert info["table_inventory_included"] is False
    assert info["table_inventory_tool"] == "list_tables"
    assert "Charts" not in info["used_ranges"]


def test_error_inventory_native_read_recovery_keeps_source_and_table_bounds(source):
    before, mtime = source.read_bytes(), source.stat().st_mtime_ns
    failed = json.loads(server.read_excel_as_table(str(source), "Sales"))
    assert not failed["ok"] and failed["error"]["code"] == "sheet_name_is_table"
    inventory = json.loads(
        server.list_tables(
            str(source), **failed["error"]["details"]["suggested_arguments"]
        )
    )
    identity = inventory["data"]["tables"][0]
    assert identity["table_name"] == "Sales" and identity["sheet_name"] == "Data"
    result = json.loads(
        server.read_excel_table(str(source), table_name=identity["table_name"])
    )
    assert result["ok"] and result["data"]["range"] == "A3:B5"
    assert result["data"]["rows"] == [[1, 10], [2, 20]]
    assert source.read_bytes() == before and source.stat().st_mtime_ns == mtime


@pytest.mark.parametrize(
    "reader",
    [
        lambda p: tables.list_excel_tables(p, sheet_name="Sales"),
        lambda p: tables.read_excel_table(p, "Sales", sheet_name="Sales"),
    ],
)
def test_bad_sheet_filter_does_not_loop_on_same_inventory_filter(source, reader):
    with pytest.raises(DataError) as caught:
        reader(str(source))
    assert caught.value.code == "sheet_name_is_table"
    assert caught.value.details["suggested_arguments"] == {"sheet_name": None}


@pytest.mark.parametrize(
    "name,code",
    [
        ("Data", "table_name_is_worksheet"),
        ("Charts", "table_name_is_chartsheet"),
        ("Missing", "table_not_found"),
    ],
)
def test_table_names_never_fall_back_to_a_sheet_or_first_table(source, name, code):
    before = source.read_bytes()
    failed = json.loads(server.read_excel_table(str(source), name))
    assert not failed["ok"] and failed["error"]["code"] == code
    assert failed["error"]["suggested_next_tool"] == "list_tables"
    assert not failed["error"]["details"]["automatic_selection"]
    assert "replace_with" not in json.dumps(failed)
    assert source.read_bytes() == before


def test_existing_worksheet_takes_precedence_over_same_named_table_elsewhere(source):
    wb = load_workbook(source)
    wb.create_sheet("Sales")["A1"] = "This is the actual worksheet"
    wb.save(source)
    wb.close()
    result = data.read_excel_range(str(source), "Sales", "A1", "A1")
    assert result == [["This is the actual worksheet"]]
    assert tables.read_excel_table(str(source), "Sales")["rows"] == [[1, 10], [2, 20]]


def test_exact_case_only_not_fuzzy_or_trimmed(source):
    for name in ("sales", " Sales", "Sales "):
        failed = json.loads(server.quick_read(str(source), sheet_name=name))
        assert failed["error"]["code"] == "worksheet_not_found"
        assert failed["error"]["details"]["table_match_count"] == 0
        assert failed["error"]["details"]["table_matches"] == []


def test_mutation_error_does_not_redirect_to_table_worksheet(source):
    before, mtime = source.read_bytes(), source.stat().st_mtime_ns
    result = json.loads(server.format_range(str(source), "Sales", "A1", bold=True))
    assert not result["ok"] and "not found" in result["error"]["message"]
    assert source.read_bytes() == before and source.stat().st_mtime_ns == mtime


def test_table_creation_preserves_selection_error_without_mutation(source):
    before, mtime = source.read_bytes(), source.stat().st_mtime_ns
    result = json.loads(server.create_table(str(source), "Sales", "A1:B3", "NewTable"))
    assert not result["ok"] and result["error"]["type"] == "DataError"
    assert result["error"]["code"] == "sheet_name_is_table"
    assert source.read_bytes() == before and source.stat().st_mtime_ns == mtime


def test_legitimate_creation_can_share_name_with_a_table(source):
    result = json.loads(server.create_worksheet(str(source), "Sales"))
    assert result["ok"]
    wb = load_workbook(source)
    assert "Sales" in wb.sheetnames and "Sales" in wb["Data"].tables
    wb.close()


def test_read_only_workbook_without_table_metadata_still_has_clean_error(source):
    wb = load_workbook(source, read_only=True)
    try:
        with pytest.raises(WorkbookError) as caught:
            workbook.require_worksheet(wb, "Missing")
        assert caught.value.code == "worksheet_not_found"
    finally:
        wb.close()


def test_read_only_worksheet_is_not_mislabeled_a_chartsheet(source):
    wb = load_workbook(source, read_only=True)
    try:
        with pytest.raises(DataError) as caught:
            tables._find_table(wb, "Data")
        assert caught.value.code == "table_name_is_worksheet"
    finally:
        wb.close()


def test_oversized_scope_diagnostic_is_bounded():
    wb = Workbook()
    try:
        with pytest.warns(UserWarning):
            wb.active.title = "\u00e4" * 1500
        with pytest.raises(DataError) as caught:
            tables._find_table(wb, "Missing", sheet_name=wb.active.title)
        assert caught.value.details["scope_omitted"]
        assert caught.value.details["scope_sheet"] is None
        assert len(server._error_response("test", caught.value)) < 3000
    finally:
        wb.close()


def test_custom_error_class_is_preserved():
    wb = Workbook()
    try:
        with pytest.raises(ValueError) as caught:
            workbook.require_worksheet(wb, "Missing", error_cls=ValueError)
        assert type(caught.value) is ValueError
        assert caught.value.code == "worksheet_not_found"
    finally:
        wb.close()


def test_diagnostic_matches_are_bounded_and_never_selected():
    wb = Workbook()
    try:
        for index in range(7):
            ws = wb.create_sheet(f"Data{index}")
            # Deliberately invalid duplicate identities in memory only: diagnostics
            # must not resolve them by choosing the first candidate.
            ws._tables["Sales"] = Table(displayName="Sales", ref="A1:B2")
        error = workbook.worksheet_not_found_error(wb, "Sales")
        assert error.details["table_match_count"] == 7
        assert len(error.details["table_matches"]) == 3
        assert error.details["matches_truncated"]
        assert not error.details["automatic_selection"]
    finally:
        wb.close()


def test_oversized_identifier_is_not_returned_as_a_partial_identity():
    wb = Workbook()
    name = "\u00e4" * 1500
    try:
        wb.active._tables[name] = Table(displayName=name, ref="A1:B2")
        error = workbook.worksheet_not_found_error(wb, name)
        assert error.details["table_match_count"] == 1
        assert (
            error.details["table_matches"] == [] and error.details["matches_truncated"]
        )
        wire = server._error_response("test", error)
        assert len(wire) < 3000 and name not in wire
    finally:
        wb.close()


def test_incomplete_table_metadata_does_not_break_missing_sheet_error():
    wb = Workbook()
    try:
        wb.active._tables["Sales"] = Table(displayName="Sales")
        error = workbook.worksheet_not_found_error(wb, "Sales")
        assert error.code == "sheet_name_is_table"
        assert error.details["table_match_count"] == 1
        assert error.details["table_matches"] == []
        assert error.details["matches_truncated"]
    finally:
        wb.close()


def test_actual_stdio_namespace_error_and_inventory_recovery(source, tmp_path):
    before, mtime = source.read_bytes(), source.stat().st_mtime_ns
    root = Path(__file__).resolve().parents[1]
    params = StdioServerParameters(
        command=sys.executable,
        args=["-m", "excel_mcp", "stdio"],
        cwd=tmp_path,
        env={"PYTHONPATH": str(root / "src")},
    )

    async def exercise():
        with (tmp_path / "stderr.log").open("w") as log:
            async with stdio_client(params, errlog=log) as (read, write):
                async with ClientSession(read, write) as session:
                    await session.initialize()

                    async def call(name, **args):
                        response = await session.call_tool(
                            name, {"filepath": str(source), **args}
                        )
                        assert not response.isError
                        return json.loads(
                            "".join(
                                c.text for c in response.content if c.type == "text"
                            )
                        )

                    metadata = await call("get_workbook_metadata")
                    assert not metadata["data"]["table_inventory_included"]
                    failed = await call("read_excel_as_table", sheet_name="Sales")
                    assert failed["error"]["code"] == "sheet_name_is_table"
                    inventory = await call(
                        failed["error"]["suggested_next_tool"],
                        **failed["error"]["details"]["suggested_arguments"],
                    )
                    identity = inventory["data"]["tables"][0]
                    result = await call(
                        "read_excel_table", table_name=identity["table_name"]
                    )
                    assert result["data"]["rows"] == [[1, 10], [2, 20]]
                    inverse = await call("read_excel_table", table_name="Data")
                    assert inverse["error"]["code"] == "table_name_is_worksheet"

    asyncio.run(asyncio.wait_for(exercise(), timeout=30))
    assert source.read_bytes() == before and source.stat().st_mtime_ns == mtime
