"""Synthetic agent report acceptance against an independent Decimal oracle."""

import asyncio
from collections import defaultdict
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import sys
import time

from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table


HEADERS = ["LineId", "Region", "Product", "Amount", "Status", "Memo"]
OUTPUT_HEADERS = ["Region", "Product", "Revenue", "Lines", "Average"]


def _sales_rows(count, start=0):
    rows = []
    for index in range(start, start + count):
        rows.append(
            {
                "LineId": f"LINE-{index:06d}",
                "Region": ["East", "North", "South", "West"][index % 4],
                "Product": ["Gadget", "Service", "Widget"][(index // 4) % 3],
                "Amount": Decimal((index * 137) % 99999 + 10) / 100,
                "Status": "cancelled" if index % 17 == 0 else "posted",
                "Memo": "=1+1",
            }
        )
    return rows


def _oracle(rows):
    # Deliberately independent of SheetForge query, filter, grouping and checks.
    grouped = defaultdict(list)
    for row in rows:
        if row["Status"] == "posted":
            grouped[(row["Region"], row["Product"])].append(row["Amount"])
    return {
        key: (
            sum(amounts, Decimal(0)),
            len(amounts),
            sum(amounts, Decimal(0)) / len(amounts),
        )
        for key, amounts in grouped.items()
    }


def _replace_source(wb, rows, headers):
    ws = wb["Data"]
    if "Sales" in ws.tables:
        del ws.tables["Sales"]
    ws.delete_rows(3, ws.max_row)
    for col, header in enumerate(headers, start=2):
        ws.cell(3, col).value = header
    for row_number, record in enumerate(rows, start=4):
        for col, header in enumerate(headers, start=2):
            value = record[header]
            ws.cell(row_number, col).value = (
                float(value) if isinstance(value, Decimal) else value
            )
    total_row = len(rows) + 4
    ws.cell(total_row, 2).value = "TOTAL"
    amount_column = 2 + headers.index("Amount")
    from openpyxl.utils import get_column_letter

    letter = get_column_letter(amount_column)
    ws.cell(total_row, amount_column).value = f"=SUM({letter}4:{letter}{total_row - 1})"
    ws.add_table(Table(displayName="Sales", ref=f"B3:G{total_row}", totalsRowCount=1))


def _make_book(path, rows):
    wb = Workbook()
    wb.active.title = "Data"
    wb["Data"]["A1"] = "Synthetic input only"
    _replace_source(wb, rows, HEADERS)
    report = wb.create_sheet("Report")
    report["A1"] = "Unaffected report title"
    report["K2"] = "=SUM(1,2)"
    report["D4"].fill = PatternFill(fill_type="solid", fgColor="FF1F4E78")
    report.freeze_panes = "D5"
    report.print_area = "A1:U45"
    chart = BarChart()
    chart.title = "Revenue"
    chart.add_data(
        Reference(report, min_col=6, min_row=4, max_row=16), titles_from_data=True
    )
    report.add_chart(chart, "K4")
    wb.save(path)
    wb.close()


def _source_and_layout(path):
    wb = load_workbook(path)
    try:
        ws, report = wb["Data"], wb["Report"]
        return {
            "source": list(ws.values),
            "table": ws.tables["Sales"].ref,
            "table_columns": [
                (column.id, column.name) for column in ws.tables["Sales"].tableColumns
            ],
            "title": report["A1"].value,
            "formula": report["K2"].value,
            "header_fill": report["D4"].fill.fgColor.rgb,
            "freeze": report.freeze_panes,
            "print_area": str(report.print_area),
            "chart": str(report._charts[0].to_tree().attrib),
            "chart_series": report._charts[0].series[0].val.numRef.f,
            "chart_anchor": (
                report._charts[0].anchor._from.col,
                report._charts[0].anchor._from.row,
            ),
        }
    finally:
        wb.close()


def _assert_output(grid, expected):
    assert grid[0] == OUTPUT_HEADERS
    body = [row for row in grid[1:] if any(value is not None for value in row)]
    actual = {(row[0], row[1]): row[2:] for row in body}
    assert len(body) == len(actual) == len(expected)
    for key, (total, count, average) in expected.items():
        revenue, lines, mean = actual[key]
        assert abs(Decimal(str(revenue)) - total) <= Decimal("0.000000001")
        assert lines == count
        assert abs(Decimal(str(mean)) - average) <= Decimal("0.000000001")
    assert all(all(value is None for value in row) for row in grid[len(body) + 1 :])


def test_agent_report_replay_against_decimal_oracle(tmp_path):
    root = Path(__file__).resolve().parents[1]
    path = tmp_path / "synthetic-sales.xlsx"
    rows = _sales_rows(1000)
    _make_book(path, rows)
    recipe = {
        "version": 1,
        "source": {"table_name": "Sales"},
        "filters": [{"field": "Status", "op": "eq", "value": "posted"}],
        "group_by": ["Region", "Product"],
        "metrics": [
            {"op": "sum", "field": "Amount", "as": "Revenue"},
            {"op": "count", "as": "Lines"},
            {"op": "avg", "field": "Amount", "as": "Average"},
        ],
        "output": {
            "sheet_name": "Report",
            "range_ref": "D4:H40",
            "allow_overwrite": True,
        },
        "checks": [
            {"type": "unique_key", "fields": ["LineId"]},
            {"type": "row_count", "min": 1, "max": 12},
            {
                "type": "sum_matches",
                "source_field": "Amount",
                "output_field": "Revenue",
            },
        ],
    }
    parameters = StdioServerParameters(
        command=sys.executable,
        args=["-m", "excel_mcp", "stdio"],
        cwd=root,
        env={"PYTHONPATH": str(root / "src")},
    )
    receipts = []

    async def exercise():
        nonlocal rows
        with (tmp_path / "stderr.log").open("w") as stderr:
            async with stdio_client(parameters, errlog=stderr) as (read, write):
                async with ClientSession(read, write) as session:
                    await session.initialize()

                    async def call(name, args):
                        start = time.monotonic()
                        response = await session.call_tool(
                            name, {"filepath": str(path), **args}
                        )
                        assert not response.isError
                        text = response.content[0].text
                        payload = json.loads(text)
                        receipts.append(
                            {
                                "tool": name,
                                "characters": len(text),
                                "seconds": round(time.monotonic() - start, 4),
                            }
                        )
                        assert len(text) <= 50000
                        assert payload["ok"], payload
                        return payload["data"]

                    await call(
                        "describe_dataset", {"table_name": "Sales", "sample_rows": 3}
                    )
                    for phase in ("initial", "growth_and_reorder", "shrink"):
                        if phase != "initial":
                            if phase == "growth_and_reorder":
                                rows += _sales_rows(200, start=1000)
                            else:
                                rows = [row for row in rows if row["Region"] == "North"]
                            wb = load_workbook(path)
                            try:
                                _replace_source(wb, rows, list(reversed(HEADERS)))
                                wb.save(path)
                            finally:
                                wb.close()
                        expected = _oracle(rows)
                        baseline = path.read_bytes()
                        invariant = _source_and_layout(path)
                        preview = await call(
                            "run_workbook_recipe", {"recipe": recipe, "sample_limit": 3}
                        )
                        assert path.read_bytes() == baseline
                        assert (
                            preview["ready_to_commit"]
                            and preview["verification"]["passed"]
                        )
                        assert preview["source_row_count"] == len(rows)
                        assert preview["output_row_count"] == len(expected)
                        assert preview["source"]["totals_rows_excluded"] == 1
                        committed = await call(
                            "run_workbook_recipe",
                            {
                                "recipe": recipe,
                                "mode": "commit",
                                "sample_limit": 3,
                                "recipe_token": preview["recipe_token"],
                                "expected_workbook_sha256": preview[
                                    "expected_workbook_sha256"
                                ],
                            },
                        )
                        assert committed["persisted"]
                        assert (
                            committed["receipt"]["before_sha256"]
                            == hashlib.sha256(baseline).hexdigest()
                        )
                        assert (
                            Path(committed["receipt"]["snapshot"]["path"]).read_bytes()
                            == baseline
                        )
                        assert _source_and_layout(path) == invariant
                        readback = await call(
                            "read_data_from_excel",
                            {
                                "sheet_name": "Report",
                                "start_cell": "D4",
                                "end_cell": "H40",
                                "values_only": True,
                            },
                        )
                        _assert_output(readback["values"], expected)

    asyncio.run(asyncio.wait_for(exercise(), timeout=60))
    metrics = {
        "fixture": "synthetic, not customer acceptance",
        "phases": 3,
        "max_source_rows": 1200,
        "calls": receipts,
        "tool_call_count": len(receipts),
        "response_characters": sum(item["characters"] for item in receipts),
    }
    (tmp_path / "acceptance-receipt.json").write_text(json.dumps(metrics, indent=2))
    assert len(receipts) == 10
    assert metrics["response_characters"] < 100000
