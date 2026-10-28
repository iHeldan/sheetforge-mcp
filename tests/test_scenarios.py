import copy
import hashlib
import json
import os
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableFormula

from excel_mcp import changeset, explanations, scenarios, workbook
from excel_mcp.exceptions import (
    PreconditionFailedError,
    ResponseTooLargeError,
    ValidationError,
    WorkbookError,
)


@pytest.fixture
def scenario_case(tmp_path):
    path = tmp_path / "budget.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Id", "Unit", "Budget", "Note"])
    for row in [
        (1, "North", 15000, "one"),
        (2, "North", 25000, "two"),
        (3, "South", 30000, "three"),
        (4, "East", 30000, "four"),
    ]:
        ws.append(row)
    ws.add_table(Table(displayName="Budget", ref="A1:D5"))
    ws["F1"] = "=1+1"
    ws["F1"].font = Font(bold=True)
    wb.save(path)
    wb.close()
    spec = {
        "version": 1,
        "recipe": {
            "version": 1,
            "source": {"table_name": "Budget"},
            "group_by": ["Unit"],
            "metrics": [
                {"op": "sum", "field": "Budget", "as": "Total"},
                {"op": "count", "as": "Projects"},
            ],
            "output": {"sheet_name": "UnusedReport", "range_ref": "A1:C20"},
        },
        "key_fields": ["Id"],
        "guards": [
            {"metric": "Total", "scope": "overall", "max": 100000},
            {"metric": "Total", "scope": "each_group", "min": 10000},
        ],
        "scenarios": [
            {"id": "A", "updates": [{"key": {"Id": 1}, "values": {"Budget": 18000}}]},
            {
                "id": "B",
                "updates": [
                    {"key": {"Id": 1}, "values": {"Budget": 20000}},
                    {"key": {"Id": 3}, "values": {"Budget": 25000}},
                ],
            },
            {
                "id": "C",
                "updates": [
                    {"key": {"Id": 4}, "values": {"Budget": 9500}},
                    {"key": {"Id": 2}, "values": {"Budget": 43500}},
                ],
            },
        ],
    }
    return path, spec


def compare(case, **kwargs):
    path, spec = case
    return scenarios.evaluate_workbook_scenarios(str(path), spec, **kwargs)


def commit(case, preview, selected="B", **kwargs):
    item = next(s for s in preview["scenarios"] if s["id"] == selected)
    return compare(
        case,
        mode="commit",
        selected_scenario=selected,
        expected_workbook_sha256=preview["expected_workbook_sha256"],
        scenario_token=item["scenario_token"],
        **kwargs,
    )


def edit(path, cell, value):
    wb = load_workbook(path)
    wb["Data"][cell] = value
    wb.save(path)
    wb.close()


def test_compare_independent_alternatives_and_commit_with_readback(scenario_case):
    path, spec = scenario_case
    original, mtime = path.read_bytes(), path.stat().st_mtime_ns
    frozen_spec = copy.deepcopy(spec)
    result = compare(scenario_case)
    assert path.read_bytes() == original and path.stat().st_mtime_ns == mtime
    assert spec == frozen_spec
    assert [s["passes_guards"] for s in result["scenarios"]] == [False, True, False]
    assert [s["ready_to_commit"] for s in result["scenarios"]] == [False, True, False]
    assert [s["metrics"][0]["before"] for s in result["scenarios"]] == ["100000"] * 3
    assert [s["metrics"][0]["after"] for s in result["scenarios"]] == [
        "103000",
        "100000",
        "98000",
    ]
    assert not list(path.parent.glob("*.snapshot*.xlsx"))
    saved = commit(scenario_case, result)
    assert (
        saved["persisted"]
        and saved["assertions_passed"] == saved["receipt"]["assertion_count"]
    )
    wb = load_workbook(path)
    assert [wb["Data"].cell(r, 3).value for r in range(2, 6)] == [
        20000,
        25000,
        25000,
        30000,
    ]
    assert [wb["Data"].cell(r, 4).value for r in range(2, 6)] == [
        "one",
        "two",
        "three",
        "four",
    ]
    assert wb["Data"]["F1"].value == "=1+1" and wb["Data"]["F1"].font.bold
    assert wb.sheetnames == ["Data"] and wb["Data"].tables["Budget"].ref == "A1:D5"
    wb.close()
    assert Path(saved["receipt"]["snapshot"]["path"]).read_bytes() == original


@pytest.mark.parametrize("selected", ["A", "C"])
def test_failed_guards_cannot_commit_any_cells(scenario_case, selected):
    preview = compare(scenario_case)
    original = scenario_case[0].read_bytes()
    with pytest.raises(PreconditionFailedError, match="violates guards"):
        commit(scenario_case, preview, selected)
    assert scenario_case[0].read_bytes() == original


def test_unpatched_row_invalidates_global_approval(scenario_case):
    path, _ = scenario_case
    first = compare(scenario_case)
    edit(path, "C5", 31000)
    original = path.read_bytes()
    with pytest.raises(PreconditionFailedError, match="changed"):
        commit(scenario_case, first)
    fresh = compare(scenario_case)
    assert not fresh["scenarios"][1]["passes_guards"]
    assert fresh["scenarios"][1]["metrics"][0]["after"] == "101000"
    with pytest.raises(PreconditionFailedError, match="violates guards"):
        commit(scenario_case, fresh)
    assert path.read_bytes() == original


def test_benign_save_requires_fresh_approval_and_is_preserved(scenario_case):
    path, _ = scenario_case
    first = compare(scenario_case)
    edit(path, "D5", "human review")
    with pytest.raises(PreconditionFailedError):
        commit(scenario_case, first)
    fresh = compare(scenario_case)
    assert (
        fresh["scenarios"][1]["scenario_token"]
        != first["scenarios"][1]["scenario_token"]
    )
    commit(scenario_case, fresh)
    wb = load_workbook(path)
    assert wb["Data"]["D5"].value == "human review"
    wb.close()


@pytest.mark.parametrize(
    "mutation", ["guards", "other_alternative", "recipe", "path", "selection"]
)
def test_full_contract_and_path_bound_to_token(scenario_case, tmp_path, mutation):
    path, spec = scenario_case
    preview = compare(scenario_case)
    selected = "B"
    if mutation == "guards":
        spec["guards"][0]["max"] = 200000
    elif mutation == "other_alternative":
        spec["scenarios"][0]["updates"][0]["values"]["Budget"] = 19000
    elif mutation == "recipe":
        spec["recipe"]["output"]["sheet_name"] = "Elsewhere"
    elif mutation == "path":
        other = tmp_path / "identical.xlsx"
        other.write_bytes(path.read_bytes())
        scenario_case = other, spec
    else:
        preview["scenarios"][0]["scenario_token"] = preview["scenarios"][1][
            "scenario_token"
        ]
        selected = "A"
    original = scenario_case[0].read_bytes()
    with pytest.raises(PreconditionFailedError, match="token"):
        commit(scenario_case, preview, selected)
    assert scenario_case[0].read_bytes() == original


def test_sample_limit_is_not_authority(scenario_case):
    one = compare(scenario_case, sample_limit=1)
    five = compare(scenario_case, sample_limit=5)
    assert [s["scenario_token"] for s in one["scenarios"]] == [
        s["scenario_token"] for s in five["scenarios"]
    ]
    assert commit(scenario_case, one, sample_limit=5)["persisted"]


def test_noop_observation_does_not_save_or_create_snapshot(scenario_case):
    path, spec = scenario_case
    spec["scenarios"] = [
        {"id": "B", "updates": [{"key": {"Id": 1}, "values": {"Budget": 15000}}]}
    ]
    original, mtime = path.read_bytes(), path.stat().st_mtime_ns
    preview = compare(scenario_case)
    result = commit(scenario_case, preview)
    assert not result["persisted"] and result["scenario"]["changed_cell_count"] == 0
    assert path.read_bytes() == original and path.stat().st_mtime_ns == mtime
    assert not list(path.parent.glob("*.snapshot*.xlsx"))


def test_disappearing_group_fails_instead_of_evading_minimum(scenario_case):
    _, spec = scenario_case
    spec["scenarios"] = [
        {"id": "B", "updates": [{"key": {"Id": 4}, "values": {"Unit": "North"}}]}
    ]
    item = compare(scenario_case)["scenarios"][0]
    assert not item["passes_guards"]
    check = item["guards"][1]
    assert check["evaluated_count"] == 3 and check["failed_count"] == 1
    assert check["failures_sample"][0]["reason"] == "missing_group"
    assert check["failures_sample"][0]["group"]["Unit"]["value"] == "East"


def test_empty_group_coverage_never_passes_vacuously(scenario_case):
    _, spec = scenario_case
    spec["recipe"]["filters"] = [{"field": "Budget", "op": "gt", "value": 1000000}]
    item = compare(scenario_case)["scenarios"][1]
    assert item["metrics"][0]["after"] == "0"
    assert not item["passes_guards"] and item["guards"][1]["empty_coverage"]


def test_new_groups_also_checked_and_count_guards(scenario_case):
    _, spec = scenario_case
    spec["guards"] = [{"metric": "Projects", "scope": "each_group", "min": 2}]
    spec["scenarios"] = [
        {"id": "B", "updates": [{"key": {"Id": 1}, "values": {"Unit": "New"}}]}
    ]
    check = compare(scenario_case)["scenarios"][0]["guards"][0]
    assert check["evaluated_count"] == 4 and check["failed_count"] == 4


def test_exact_fraction_guard_not_relaxed_by_recipe_tolerance(scenario_case):
    path, spec = scenario_case
    for r, value in enumerate([0.1, 0.2, 0, 0], 2):
        edit(path, f"C{r}", value)
    spec["guards"] = [{"metric": "Total", "scope": "overall", "max": 0.3}]
    spec["scenarios"] = [
        {
            "id": "B",
            "updates": [{"key": {"Id": 3}, "values": {"Budget": 0.000000000001}}],
        }
    ]
    item = compare(scenario_case)["scenarios"][0]
    assert item["metrics"][0]["after"] == "0.300000000001"
    assert not item["passes_guards"]
    spec["guards"][0]["max"] = 0.300000000001
    assert compare(scenario_case)["scenarios"][0]["passes_guards"]


def test_integer_valued_float_roundtrip_and_null_are_supported(scenario_case):
    _, spec = scenario_case
    spec["guards"] = [{"metric": "Total", "scope": "overall", "max": 100000}]
    spec["scenarios"] = [
        {
            "id": "B",
            "updates": [
                {"key": {"Id": 1}, "values": {"Budget": 20000.0}},
                {"key": {"Id": 3}, "values": {"Budget": None}},
            ],
        }
    ]
    preview = compare(scenario_case)
    assert preview["scenarios"][0]["metrics"][0]["after"] == "75000"
    assert commit(scenario_case, preview)["persisted"]


@pytest.mark.parametrize(
    "bounds",
    [
        {"min": True},
        {"max": None},
        {"min": "10"},
        {"min": 2, "max": 1},
        {"max": float("inf")},
    ],
)
def test_invalid_bounds_fail_before_capture(scenario_case, monkeypatch, bounds):
    scenario_case[1]["guards"] = [{"metric": "Total", "scope": "overall", **bounds}]
    monkeypatch.setattr(
        changeset, "_stage_workbook", lambda *a: pytest.fail("must not capture")
    )
    with pytest.raises(ValidationError):
        compare(scenario_case)


@pytest.mark.parametrize(
    "mutation",
    [
        "force",
        "version",
        "empty_guards",
        "duplicate_id",
        "duplicate_update",
        "empty_updates",
        "key_write",
        "formula_output",
        "unsupported_metric",
        "unknown_metric",
        "unknown_scope",
        "invalid_unicode",
    ],
)
def test_closed_contract(scenario_case, mutation):
    path, spec = scenario_case
    if mutation == "force":
        spec["force"] = True
    elif mutation == "version":
        spec["version"] = True
    elif mutation == "empty_guards":
        spec["guards"] = []
    elif mutation == "duplicate_id":
        spec["scenarios"][0]["id"] = "B"
    elif mutation == "duplicate_update":
        spec["scenarios"][0]["updates"] *= 2
    elif mutation == "empty_updates":
        spec["scenarios"][0]["updates"] = []
    elif mutation == "key_write":
        spec["scenarios"][0]["updates"][0]["values"] = {"Id": 9}
    elif mutation == "formula_output":
        spec["scenarios"][0]["updates"][0]["values"] = {"Budget": "=1+1"}
    elif mutation == "unsupported_metric":
        spec["recipe"]["metrics"][0]["op"] = "avg"
    elif mutation == "unknown_metric":
        spec["guards"][0]["metric"] = "Typo"
    elif mutation == "unknown_scope":
        spec["guards"][0]["scope"] = []
    else:
        spec["scenarios"][0]["updates"][0]["values"]["Note"] = "\ud800"
    original = path.read_bytes()
    with pytest.raises(ValidationError):
        compare(scenario_case)
    assert path.read_bytes() == original


@pytest.mark.parametrize(
    "fault", ["duplicate", "formula", "calculated", "protected", "merged"]
)
def test_bad_native_table_inputs_are_not_partially_verified(scenario_case, fault):
    path, _ = scenario_case
    wb = load_workbook(path)
    ws = wb["Data"]
    if fault == "duplicate":
        ws["A5"] = 1
    elif fault == "formula":
        ws["C5"] = "=SUM(C2:C4)"
    elif fault == "calculated":
        ws.tables["Budget"].tableColumns[2].calculatedColumnFormula = TableFormula(
            attr_text="1+1"
        )
    elif fault == "protected":
        ws.protection.sheet = True
    else:
        ws.merge_cells("C2:D2")
    wb.save(path)
    wb.close()
    original = path.read_bytes()
    with pytest.raises(ValidationError):
        compare(scenario_case)
    assert path.read_bytes() == original


def test_race_between_capture_and_changeset_is_rejected(scenario_case, monkeypatch):
    path, _ = scenario_case
    real = changeset.apply_workbook_changeset

    def race(*args, **kwargs):
        edit(path, "D2", "concurrent")
        return real(*args, **kwargs)

    monkeypatch.setattr(changeset, "apply_workbook_changeset", race)
    with pytest.raises(PreconditionFailedError):
        compare(scenario_case)
    wb = load_workbook(path)
    assert wb["Data"]["C2"].value == 15000 and wb["Data"]["D2"].value == "concurrent"
    wb.close()


@pytest.mark.parametrize("fault", ["guard", "cancelled", "exception"])
def test_post_replace_verification_rolls_back_exact_bytes(
    scenario_case, monkeypatch, fault
):
    path, _ = scenario_case
    preview = compare(scenario_case)
    original = path.read_bytes()
    real = workbook._verify_saved_workbook
    triggered = False

    def corrupt_once(filepath):
        nonlocal triggered
        if Path(filepath) == path and not triggered:
            triggered = True
            if fault == "exception":
                raise WorkbookError("injected reopen failure")
            wb = load_workbook(path)
            wb["Data"]["C5"] = 31000
            if fault == "cancelled":
                wb["Data"]["C3"] = 24000
            wb.save(path)
            wb.close()
        return real(filepath)

    monkeypatch.setattr(workbook, "_verify_saved_workbook", corrupt_once)
    with pytest.raises(WorkbookError):
        commit(scenario_case, preview)
    assert triggered and path.read_bytes() == original


def test_independent_evaluator_disagreement_is_not_a_pass(scenario_case, monkeypatch):
    real = explanations._groups

    def wrong(records, recipe):
        result = real(records, recipe)
        next(iter(result.values()))[0] += 1
        return result

    monkeypatch.setattr(explanations, "_groups", wrong)
    original = scenario_case[0].read_bytes()
    with pytest.raises(PreconditionFailedError, match="independent"):
        compare(scenario_case)
    assert scenario_case[0].read_bytes() == original


def test_commit_summary_budget_checked_before_mutation(scenario_case, monkeypatch):
    preview = compare(scenario_case)
    original = scenario_case[0].read_bytes()
    monkeypatch.setattr(scenarios, "MAX_COMMIT_SUMMARY_CHARS", 1)
    monkeypatch.setattr(
        changeset, "apply_workbook_changeset", lambda *a, **k: pytest.fail("no write")
    )
    with pytest.raises(ResponseTooLargeError):
        commit(scenario_case, preview)
    assert scenario_case[0].read_bytes() == original


def test_non_regular_input_rejected_before_open(tmp_path, scenario_case):
    directory = tmp_path / "directory.xlsx"
    directory.mkdir()
    with pytest.raises(ValidationError, match="regular"):
        compare((directory, scenario_case[1]))


def test_direct_changeset_scenario_assertion_cannot_skip_guards(scenario_case):
    path, spec = scenario_case
    assertion = {
        "type": "table_scenario_matches",
        "recipe": spec["recipe"],
        "key_fields": spec["key_fields"],
        "guards": spec["guards"],
        "expected_result_sha256": "0" * 64,
    }
    ops = [
        {
            "tool": "write_data_to_excel",
            "args": {"sheet_name": "Data", "start_cell": "C2", "data": [[18000]]},
        }
    ]
    original = path.read_bytes()
    result = changeset.apply_workbook_changeset(str(path), ops, [assertion])
    assert (
        not result["ready_to_commit"]
        and not result["assertions"][0]["details"]["passes_guards"]
    )
    with pytest.raises(PreconditionFailedError):
        changeset.apply_workbook_changeset(
            str(path),
            ops,
            [assertion],
            mode="commit",
            expected_workbook_sha256=result["expected_workbook_sha256"],
            changeset_token=result["changeset_token"],
        )
    assert path.read_bytes() == original


def test_compare_pin_and_missing_commit_credentials(scenario_case):
    with pytest.raises(PreconditionFailedError):
        compare(scenario_case, expected_workbook_sha256="0" * 64)
    with pytest.raises(ValidationError):
        compare(scenario_case, mode="commit", selected_scenario="B")
    result = compare(scenario_case)
    assert (
        result["expected_workbook_sha256"]
        == hashlib.sha256(scenario_case[0].read_bytes()).hexdigest()
    )


@pytest.mark.parametrize("sample", [0, 6, True, None, []])
def test_invalid_sample_size(scenario_case, sample):
    with pytest.raises(ValidationError):
        compare(scenario_case, sample_limit=sample)


@pytest.mark.parametrize("mode", ["preview", "force", None, []])
def test_invalid_mode(scenario_case, mode):
    with pytest.raises(ValidationError):
        compare(scenario_case, mode=mode)


def test_mcp_domain_errors_and_path_boundary(scenario_case, monkeypatch):
    import excel_mcp.server as server

    path, spec = scenario_case
    result = json.loads(server.evaluate_workbook_scenarios(str(path), spec))
    assert result["ok"] and result["data"]["scenarios"][1]["ready_to_commit"]

    def deny(_):
        raise ValidationError("outside allowed root")

    monkeypatch.setattr(server, "get_excel_path", deny)
    monkeypatch.setattr(
        server,
        "evaluate_workbook_scenarios_impl",
        lambda *a, **kw: pytest.fail("no domain IO"),
    )
    result = json.loads(server.evaluate_workbook_scenarios(str(path), spec))
    assert not result["ok"]


def test_fifo_rejected_before_open(scenario_case, tmp_path):
    path = tmp_path / "fifo.xlsx"
    os.mkfifo(path)
    with pytest.raises(ValidationError, match="regular"):
        compare((path, scenario_case[1]))


def test_composite_keys_offset_table_totals_and_column_order(scenario_case):
    path, spec = scenario_case
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row, values in enumerate(
        [
            ["Budget", "Note", "Unit", "Id"],
            [30000, "four", "East", 4],
            [30000, "three", "South", 3],
            [25000, "two", "North", 2],
            [15000, "one", "North", 1],
            ["=SUM(C6:C9)", "totals", None, None],
        ],
        5,
    ):
        for col, value in enumerate(values, 3):
            ws.cell(row, col).value = value
    ws.add_table(Table(displayName="Budget", ref="C5:F10", totalsRowCount=1))
    wb.save(path)
    wb.close()
    spec["key_fields"] = ["Id", "Note"]
    labels = {1: "one", 2: "two", 3: "three", 4: "four"}
    for alt in spec["scenarios"]:
        for update in alt["updates"]:
            update["key"]["Note"] = labels[update["key"]["Id"]]
    result = compare(scenario_case)
    assert result["scenarios"][1]["metrics"][0]["after"] == "100000"
    commit(scenario_case, result)
    wb = load_workbook(path)
    assert [wb["Data"].cell(r, 3).value for r in range(6, 10)] == [
        30000,
        25000,
        25000,
        20000,
    ]
    assert wb["Data"]["C10"].value == "=SUM(C6:C9)"
    assert wb["Data"].tables["Budget"].ref == "C5:F10"
    wb.close()


def test_global_only_recipe_and_baseline_guard_violation_can_be_repaired(scenario_case):
    path, spec = scenario_case
    edit(path, "C2", 20000)
    spec["recipe"]["group_by"] = []
    spec["guards"] = [{"metric": "Total", "scope": "overall", "max": 100000}]
    spec["scenarios"] = [
        {"id": "B", "updates": [{"key": {"Id": 1}, "values": {"Budget": 15000}}]}
    ]
    result = compare(scenario_case)
    assert result["scenarios"][0]["metrics"][0] == {
        "name": "Total",
        "before": "105000",
        "after": "100000",
        "delta": "-5000",
    }
    assert commit(scenario_case, result)["persisted"]


def test_unicode_commit_budget_uses_escaped_wire_size(scenario_case, monkeypatch):
    _, spec = scenario_case
    spec["scenarios"] = [
        {"id": "B", "updates": [{"key": {"Id": 1}, "values": {"Note": "\u2603" * 500}}]}
    ]
    result = compare(scenario_case)
    summary = result["scenarios"][0]
    raw = len(json.dumps(summary, ensure_ascii=False))
    escaped = len(json.dumps(summary, ensure_ascii=True))
    assert escaped > raw + 2000
    monkeypatch.setattr(scenarios, "MAX_COMMIT_SUMMARY_CHARS", raw + 1500)
    original = scenario_case[0].read_bytes()
    with pytest.raises(ResponseTooLargeError):
        commit(scenario_case, result)
    assert scenario_case[0].read_bytes() == original


@pytest.mark.parametrize(
    "limit", ["alternatives", "guards", "metrics", "spec_bytes", "targets"]
)
def test_resource_bounds_precede_capture(scenario_case, monkeypatch, limit):
    _, spec = scenario_case
    if limit == "alternatives":
        spec["scenarios"] *= 2
    elif limit == "guards":
        spec["guards"] *= 11
    elif limit == "metrics":
        spec["recipe"]["metrics"] = [{"op": "count", "as": f"M{i}"} for i in range(9)]
    elif limit == "spec_bytes":
        monkeypatch.setattr(scenarios, "MAX_SPEC_BYTES", 10)
    else:
        spec["scenarios"][0]["updates"] = [
            {"key": {"Id": i}, "values": {"Budget": 1}} for i in range(1001)
        ]
    monkeypatch.setattr(changeset, "_stage_workbook", lambda *a: pytest.fail("no IO"))
    with pytest.raises(ValidationError):
        compare(scenario_case)


def test_span_limit_reserves_the_outcome_assertion(scenario_case):
    path, spec = scenario_case
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Id", "Unit", "Budget", "Note"])
    for i in range(100):
        ws.append([i, "North", 100, "old"])
    ws.add_table(Table(displayName="Budget", ref="A1:D101"))
    wb.save(path)
    wb.close()
    spec["scenarios"] = [
        {
            "id": "B",
            "updates": [
                {"key": {"Id": i}, "values": {"Note": "new"}} for i in range(0, 100, 2)
            ],
        }
    ]
    original = path.read_bytes()
    with pytest.raises(ValidationError, match="49"):
        compare(scenario_case)
    assert path.read_bytes() == original


def test_recipe_engine_drift_invalidates_prior_token(scenario_case, monkeypatch):
    result = compare(scenario_case)
    monkeypatch.setattr(scenarios.recipes, "RECIPE_ENGINE_VERSION", 999)
    with pytest.raises(PreconditionFailedError, match="token"):
        commit(scenario_case, result)


def test_replaying_successful_commit_is_stale_not_exactly_once(scenario_case):
    first = compare(scenario_case)
    assert commit(scenario_case, first)["persisted"]
    with pytest.raises(PreconditionFailedError, match="changed"):
        commit(scenario_case, first)
    fresh = compare(scenario_case)
    assert not commit(scenario_case, fresh)["persisted"]


def test_packaged_example_matches_fixture(scenario_case):
    path, _ = scenario_case
    example = (
        Path(__file__).resolve().parents[1] / "docs/examples/budget.scenarios.json"
    )
    result = compare((path, json.loads(example.read_text())))
    assert [s["ready_to_commit"] for s in result["scenarios"]] == [False, True, False]


def test_numeric_match_covers_baseline_not_only_candidate(scenario_case):
    path, spec = scenario_case
    for row, value in enumerate([0.1234567890123456, 1, 0, 0], 2):
        edit(path, f"C{row}", value)
    spec["guards"] = [{"metric": "Total", "scope": "overall", "max": 10}]
    spec["scenarios"] = [
        {"id": "B", "updates": [{"key": {"Id": 1}, "values": {"Budget": 0.5}}]}
    ]
    item = compare(scenario_case)["scenarios"][0]
    assert item["numeric_match"] == "within_declared_tolerance"
    assert item["metrics"][0]["after"] == "1.5"


def test_core_guard_evidence_is_not_truncated_like_source_samples(scenario_case):
    _, spec = scenario_case
    spec["guards"] = [{"metric": "Total", "scope": "overall", "min": -(10**2200)}]
    spec["scenarios"] = [spec["scenarios"][1]]
    item = compare(scenario_case)["scenarios"][0]
    assert item["guards"][0]["min"] == str(-(10**2200))


def test_snapshot_conflict_is_actionable_and_never_overwritten(scenario_case):
    result = compare(scenario_case)
    snapshot = Path(result["scenarios"][1]["snapshot"]["path"])
    snapshot.write_bytes(b"owned by somebody else")
    original = scenario_case[0].read_bytes()
    fresh = compare(scenario_case)
    item = fresh["scenarios"][1]
    assert item["passes_guards"] and not item["ready_to_commit"]
    assert item["snapshot"]["status"] == "conflict"
    assert item["snapshot"]["path"] == str(snapshot)
    assert item["snapshot"]["reason"]
    with pytest.raises(PreconditionFailedError):
        commit(scenario_case, fresh)
    assert scenario_case[0].read_bytes() == original
    assert snapshot.read_bytes() == b"owned by somebody else"
