import asyncio
import copy
import hashlib
import json
from pathlib import Path
import sys

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

import excel_mcp.recipes as recipes
from excel_mcp.recipes import (
    _aggregate,
    _load_table,
    _normalize_recipe,
    run_workbook_recipe,
)
from excel_mcp.exceptions import (
    PreconditionFailedError,
    ResponseTooLargeError,
    ValidationError,
    WorkbookError,
)


@pytest.fixture
def recipe():
    return {
        "version": 1,
        "source": {"table_name": "Sales"},
        "filters": [{"field": "Amount", "op": "gte", "value": 0}],
        "group_by": ["Region"],
        "metrics": [{"op": "sum", "field": "Amount", "as": "Revenue"}],
        "output": {"sheet_name": "Report", "range_ref": "A1:B10"},
    }


@pytest.fixture
def source_book(tmp_path):
    path = tmp_path / "sales.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in [
        ["Id", "Region", "Amount"],
        [1, "North", 10],
        [2, "South", 20],
        [3, "North", 5],
    ]:
        ws.append(row)
    ws.add_table(Table(displayName="Sales", ref="A1:C4"))
    wb.save(path)
    wb.close()
    return path


def test_recipe_aggregates_without_model_data_roundtrip(source_book, recipe):
    from openpyxl import load_workbook

    normalized = _normalize_recipe(recipe)
    wb = load_workbook(source_book)
    try:
        source = _load_table(wb, normalized)
        result = _aggregate(source, normalized)
    finally:
        wb.close()
    assert result["rows"] == [["North", 15], ["South", 20]]
    assert source["bindings"] == {"Region": "B1", "Amount": "C1"}
    assert "checks" not in recipe


def test_recipe_invariance_probes_do_not_hide_cancellation_error(source_book, recipe):
    def cancellation(wb):
        for row, value in enumerate([1e16, 1, -1e16], start=2):
            wb["Data"].cell(row, 3).value = value

    _edit(source_book, cancellation)
    recipe["filters"] = []
    recipe["group_by"] = []
    preview = run_workbook_recipe(str(source_book), recipe)
    assert preview["ready_to_commit"]
    assert preview["verification"]["passed"]
    assert preview["rows_sample"] == [[1]]
    assert _commit(source_book, recipe, preview)["persisted"]


@pytest.mark.parametrize(
    "update",
    [
        {"version": True},
        {"version": 2},
        {"code": "pass"},
        {"filters": {"field": "Id"}},
        {"metrics": []},
        {"group_by": ["Region", "Region"]},
        {"checks": [{"type": "unknown"}]},
        {"output": {"sheet_name": "Report", "range_ref": "A1:XFD1048576"}},
        {"output": {"sheet_name": "Report", "range_ref": "B2:A1"}},
    ],
)
def test_recipe_rejects_unsupported_contract(recipe, update):
    candidate = copy.deepcopy(recipe)
    candidate.update(update)
    with pytest.raises(ValidationError):
        _normalize_recipe(candidate)


def _commit(path, recipe, preview):
    return run_workbook_recipe(
        str(path),
        recipe,
        mode="commit",
        recipe_token=preview["recipe_token"],
        expected_workbook_sha256=preview["expected_workbook_sha256"],
    )


def _edit(path, action):
    wb = load_workbook(path)
    try:
        action(wb)
        wb.save(path)
    finally:
        wb.close()


def _assert_no_work_artifacts(path):
    assert list(path.parent.glob(".*.sheetforge-*")) == []


def test_recipe_cleanup_failure_does_not_mask_success(
    source_book, recipe, monkeypatch, caplog
):
    preview = run_workbook_recipe(str(source_book), recipe)
    original = recipes.shutil.rmtree

    def fail_recipe_cleanup(path, *args, **kwargs):
        original(path, *args, **kwargs)
        if Path(path).name.startswith("sheetforge-recipe-"):
            raise OSError("simulated cleanup failure")

    monkeypatch.setattr(recipes.shutil, "rmtree", fail_recipe_cleanup)
    result = _commit(source_book, recipe, preview)
    assert result["persisted"]
    assert "Could not remove recipe temporary workspace" in caplog.text
    wb = load_workbook(source_book)
    try:
        assert wb["Report"]["B2"].value == 15
    finally:
        wb.close()


def test_documented_recipe_runs_unchanged(source_book):
    example = (
        Path(__file__).resolve().parents[1] / "docs/examples/regional-sales.recipe.json"
    )
    recipe = json.loads(example.read_text())
    preview = run_workbook_recipe(str(source_book), recipe)
    assert preview["verification"]["passed"]
    result = _commit(source_book, recipe, preview)
    assert result["persisted"]
    wb = load_workbook(source_book)
    try:
        assert list(wb["Report"].values)[:3] == [
            ("Region", "Revenue"),
            ("North", 15),
            ("South", 20),
        ]
    finally:
        wb.close()


def test_exact_headers_take_precedence_over_trimmed_schema_aliases(source_book, recipe):
    def change_headers(wb):
        ws = wb["Data"]
        ws["A1"] = "region!"
        ws["B1"] = " region "
        ws["C1"] = " Amount Total "
        for column, name in zip(
            ws.tables["Sales"].tableColumns, ["region!", " region ", " Amount Total "]
        ):
            column.name = name

    _edit(source_book, change_headers)
    recipe["group_by"] = [" region "]
    recipe["filters"][0]["field"] = " Amount Total "
    recipe["metrics"][0]["field"] = " Amount Total "
    preview = run_workbook_recipe(str(source_book), recipe)
    assert preview["ready_to_commit"]
    assert preview["headers"] == [" region ", "Revenue"]
    assert preview["rows_sample"] == [["North", 15], ["South", 20]]
    assert _commit(source_book, recipe, preview)["persisted"]


def test_recipe_large_source_stays_in_server(source_book, recipe):
    def expand(wb):
        ws = wb["Data"]
        for index in range(4, 1001):
            ws.append([index, "North", 1])
        ws.tables["Sales"].ref = "A1:C1001"

    _edit(source_book, expand)
    before = source_book.read_bytes()
    preview = run_workbook_recipe(str(source_book), recipe, sample_limit=1)
    assert preview["source_row_count"] == 1000
    assert preview["rows_sample"] == [["North", 1012]]
    assert preview["verification"]["lineage_sample"][0]["source_rows"] == [2]
    assert preview["verification"]["lineage_sample"][0]["source_row_count"] == 999
    assert preview["truncated"]
    assert len(json.dumps(preview)) < 10000
    assert source_book.read_bytes() == before
    assert _commit(source_book, recipe, preview)["persisted"]


@pytest.mark.parametrize(
    "op, expected",
    [
        ("sum", 0.3),
        ("avg", 0.15),
        ("min", 0.1),
        ("max", 0.2),
        ("count", 3),
        ("count_non_null", 2),
        ("count_distinct", 2),
    ],
)
def test_all_metrics_roundtrip_with_decimal_and_null_inputs(
    source_book, recipe, op, expected
):
    def decimals(wb):
        for index, value in enumerate([0.1, 0.2, None], start=2):
            wb["Data"].cell(index, 3).value = value

    _edit(source_book, decimals)
    recipe["filters"] = []
    recipe["group_by"] = []
    recipe["metrics"] = [{"op": op, "as": "Result"}]
    if op != "count":
        recipe["metrics"][0]["field"] = "Amount"
    preview = run_workbook_recipe(str(source_book), recipe)
    assert preview["ready_to_commit"], preview["assertions"]
    assert preview["rows_sample"] == [[expected]]
    assert _commit(source_book, recipe, preview)["persisted"]
    wb = load_workbook(source_book)
    try:
        assert wb["Report"]["A2"].value == expected
    finally:
        wb.close()


def test_recipe_preview_commit_and_replay_with_smaller_result(source_book, recipe):
    recipe["output"]["allow_overwrite"] = True
    recipe["checks"] = [
        {"type": "unique_key", "fields": ["Id"]},
        {"type": "sum_matches", "source_field": "Amount", "output_field": "Revenue"},
        {"type": "row_count", "min": 1, "max": 5},
    ]
    before = source_book.read_bytes()
    preview = run_workbook_recipe(str(source_book), recipe)
    assert source_book.read_bytes() == before
    assert preview["ready_to_commit"] and not preview["persisted"]
    assert preview["rows_sample"] == [["North", 15], ["South", 20]]
    assert preview["verification"]["passed"]
    assert preview["verification"]["lineage_sample"][0]["source_rows"] == [2, 4]
    committed = _commit(source_book, recipe, preview)
    assert committed["persisted"]
    assert committed["source_sha256"] == hashlib.sha256(before).hexdigest()
    snapshot = committed["receipt"]["snapshot"]["path"]
    assert Path(snapshot).read_bytes() == before
    wb = load_workbook(source_book)
    assert wb["Report"]["B2"].value == 15
    assert wb["Report"]["B3"].value == 20
    wb.close()
    # A smaller recomputation must remove the previous South result.
    recipe["filters"] = [{"field": "Region", "op": "eq", "value": "North"}]
    preview = run_workbook_recipe(str(source_book), recipe)
    _commit(source_book, recipe, preview)
    wb = load_workbook(source_book)
    assert wb["Report"]["A3"].value is None
    assert wb["Report"]["B3"].value is None
    assert wb["Data"]["C3"].value == 20
    wb.close()
    _assert_no_work_artifacts(source_book)


def test_recipe_rejects_stale_source_and_retries(source_book, recipe):
    preview = run_workbook_recipe(str(source_book), recipe)
    _edit(source_book, lambda wb: setattr(wb["Data"]["C2"], "value", 99))
    changed = source_book.read_bytes()
    with pytest.raises(PreconditionFailedError) as exc:
        _commit(source_book, recipe, preview)
    assert exc.value.code == "stale_workbook"
    assert source_book.read_bytes() == changed
    fresh = run_workbook_recipe(str(source_book), recipe)
    _commit(source_book, recipe, fresh)
    saved = source_book.read_bytes()
    with pytest.raises(PreconditionFailedError):
        _commit(source_book, recipe, fresh)
    assert source_book.read_bytes() == saved


def test_recipe_token_binds_checks_even_when_values_are_identical(source_book, recipe):
    preview = run_workbook_recipe(str(source_book), recipe)
    recipe["checks"] = [{"type": "row_count", "min": 0, "max": 999}]
    before = source_book.read_bytes()
    with pytest.raises(PreconditionFailedError) as exc:
        _commit(source_book, recipe, preview)
    assert exc.value.code == "recipe_token_mismatch"
    assert source_book.read_bytes() == before


def test_recipe_token_cannot_be_reused_for_another_path(source_book, recipe, tmp_path):
    preview = run_workbook_recipe(str(source_book), recipe)
    other = tmp_path / "other.xlsx"
    other.write_bytes(source_book.read_bytes())
    with pytest.raises(PreconditionFailedError, match="token"):
        _commit(other, recipe, preview)


@pytest.mark.parametrize(
    "check",
    [
        {"type": "unique_key", "fields": ["Region"]},
        {"type": "row_count", "min": 3, "max": 4},
        {"type": "sum_matches", "source_field": "Id", "output_field": "Revenue"},
    ],
)
def test_recipe_failed_checks_never_mutate(source_book, recipe, check):
    recipe["checks"] = [check]
    before = source_book.read_bytes()
    preview = run_workbook_recipe(str(source_book), recipe)
    assert preview["ready_to_commit"] is False
    with pytest.raises(PreconditionFailedError) as exc:
        _commit(source_book, recipe, preview)
    assert exc.value.code == "recipe_verification_failed"
    assert source_book.read_bytes() == before
    assert not list(source_book.parent.glob("*snapshot*"))


def test_recipe_detects_source_change_between_compile_and_changeset(
    source_book, recipe, monkeypatch
):
    original = recipes.changeset.apply_workbook_changeset
    edited = []

    def intervene(*args, **kwargs):
        _edit(source_book, lambda wb: setattr(wb["Data"]["C2"], "value", 42))
        edited.append(source_book.read_bytes())
        return original(*args, **kwargs)

    preview = run_workbook_recipe(str(source_book), recipe)
    monkeypatch.setattr(recipes.changeset, "apply_workbook_changeset", intervene)
    with pytest.raises(PreconditionFailedError, match="changed"):
        _commit(source_book, recipe, preview)
    assert source_book.read_bytes() == edited[0]
    _assert_no_work_artifacts(source_book)


@pytest.mark.parametrize("value", ["=1+1", "#REF!", "text", True])
def test_recipe_rejects_formula_error_or_mixed_metric_inputs(
    source_book, recipe, value
):
    _edit(source_book, lambda wb: setattr(wb["Data"]["C2"], "value", value))
    before = source_book.read_bytes()
    with pytest.raises(ValidationError):
        run_workbook_recipe(str(source_book), recipe)
    assert source_book.read_bytes() == before


def test_recipe_rejects_duplicate_and_missing_headers(source_book, recipe):
    _edit(source_book, lambda wb: setattr(wb["Data"]["A1"], "value", "region"))
    with pytest.raises(ValidationError, match="ambiguous"):
        run_workbook_recipe(str(source_book), recipe)
    _edit(source_book, lambda wb: setattr(wb["Data"]["A1"], "value", "Id"))
    recipe["group_by"] = ["region"]
    with pytest.raises(ValidationError, match="exact header"):
        run_workbook_recipe(str(source_book), recipe)


def test_recipe_follows_reordered_columns_and_expanded_table(source_book, recipe):
    initial = run_workbook_recipe(str(source_book), recipe)

    def reorder(wb):
        ws = wb["Data"]
        for row in ws.iter_rows(min_row=1, max_row=4, min_col=1, max_col=3):
            values = [cell.value for cell in row]
            for cell, value in zip(row, reversed(values)):
                cell.value = value
        ws.append([7, "North", 4])
        ws.tables["Sales"].ref = "A1:C5"

    _edit(source_book, reorder)
    preview = run_workbook_recipe(str(source_book), recipe)
    assert initial["rows_sample"][0] == ["North", 15]
    assert preview["rows_sample"][0] == ["North", 22]
    assert preview["source"]["bindings"]["Amount"] == "A1"
    assert preview["verification"]["passed"]


def test_recipe_rejects_output_overlaps_and_implicit_overwrite(source_book, recipe):
    recipe["output"] = {
        "sheet_name": "Data",
        "range_ref": "B1:C10",
        "allow_overwrite": True,
    }
    with pytest.raises(ValidationError, match="overlaps"):
        run_workbook_recipe(str(source_book), recipe)
    recipe["output"] = {"sheet_name": "Report", "range_ref": "A1:B10"}
    _edit(
        source_book,
        lambda wb: setattr(wb.create_sheet("Report")["A1"], "value", "KEEP"),
    )
    with pytest.raises(ValidationError, match="allow_overwrite"):
        run_workbook_recipe(str(source_book), recipe)
    _edit(source_book, lambda wb: wb["Report"].merge_cells("A8:B8"))
    recipe["output"]["allow_overwrite"] = True
    with pytest.raises(ValidationError, match="merged"):
        run_workbook_recipe(str(source_book), recipe)


def test_recipe_output_limit_and_source_limit_fail_closed(
    source_book, recipe, monkeypatch
):
    recipe["output"]["range_ref"] = "A1:B2"
    with pytest.raises(ValidationError, match="does not fit"):
        run_workbook_recipe(str(source_book), recipe)
    recipe["output"]["range_ref"] = "A1:B10"
    monkeypatch.setattr(recipes, "MAX_SOURCE_CELLS", 5)
    with pytest.raises(ValidationError, match="exceeds"):
        run_workbook_recipe(str(source_book), recipe)


def test_recipe_response_budget_fails_before_mutation(source_book, recipe, monkeypatch):
    preview = run_workbook_recipe(str(source_book), recipe)
    before = source_book.read_bytes()
    monkeypatch.setattr(recipes, "MAX_SUMMARY_BYTES", 1)
    with pytest.raises(ResponseTooLargeError):
        _commit(source_book, recipe, preview)
    assert source_book.read_bytes() == before


def test_recipe_rolls_back_failed_post_replace_verification(
    source_book, recipe, monkeypatch
):
    preview = run_workbook_recipe(str(source_book), recipe)
    before = source_book.read_bytes()
    original = recipes.workbook._verify_saved_workbook

    def fail_after_replace(path):
        if str(path) == str(source_book) and source_book.read_bytes() != before:
            raise WorkbookError("injected verification failure")
        return original(path)

    monkeypatch.setattr(recipes.workbook, "_verify_saved_workbook", fail_after_replace)
    with pytest.raises(WorkbookError, match="restored"):
        _commit(source_book, recipe, preview)
    assert source_book.read_bytes() == before
    _assert_no_work_artifacts(source_book)


def test_recipe_probe_failure_blocks_commit(source_book, recipe, monkeypatch):
    original = recipes._aggregate

    def inconsistent(source, normalized):
        result = original(source, normalized)
        if source["rows"][0][2] == 5:
            result["rows"][0][1] += 1
        return result

    monkeypatch.setattr(recipes, "_aggregate", inconsistent)
    preview = run_workbook_recipe(str(source_book), recipe)
    assert preview["ready_to_commit"] is False
    assert preview["verification"]["invariance_probes"][1]["passed"] is False
    with pytest.raises(PreconditionFailedError):
        _commit(source_book, recipe, preview)


def test_recipe_output_refuses_literal_formula_conversion(source_book, recipe):
    def literal(wb):
        cell = wb["Data"]["B2"]
        cell.value = '=HYPERLINK("https://example.invalid")'
        cell.data_type = "s"

    _edit(source_book, literal)
    with pytest.raises(ValidationError, match="formula-like"):
        run_workbook_recipe(str(source_book), recipe)


def test_raw_write_clears_cells_and_reports_boolean_changes(source_book):
    from excel_mcp.data import write_data

    write_data(str(source_book), "Data", [[True]], "C2")
    preview = write_data(
        str(source_book), "Data", [[1]], "C2", dry_run=True, include_changes=True
    )
    assert preview["changed_cells"] == 1
    result = write_data(str(source_book), "Data", [[None]], "C2")
    assert result["changed_cells"] == 1
    wb = load_workbook(source_book)
    assert wb["Data"]["C2"].value is None
    wb.close()


def test_recipe_mcp_roundtrip_and_compact_response(source_book, recipe):
    from excel_mcp import server

    preview = json.loads(server.run_workbook_recipe(str(source_book), recipe))
    assert preview["ok"] is True
    assert preview["data"]["mode"] == "preview"
    assert preview["data"]["ready_to_commit"]
    result = server.run_workbook_recipe(
        str(source_book),
        recipe,
        mode="commit",
        recipe_token=preview["data"]["recipe_token"],
        expected_workbook_sha256=preview["data"]["expected_workbook_sha256"],
    )
    assert len(result) < server.MCP_RESPONSE_CHAR_LIMIT
    assert json.loads(result)["data"]["persisted"] is True


def test_recipe_stdio_protocol_preview_commit_and_stale_retry(
    source_book, recipe, tmp_path
):
    from mcp import ClientSession, StdioServerParameters
    from mcp.client.stdio import stdio_client

    root = Path(__file__).resolve().parents[1]
    parameters = StdioServerParameters(
        command=sys.executable,
        args=["-m", "excel_mcp", "stdio"],
        cwd=root,
        env={"PYTHONPATH": str(root / "src")},
    )

    async def exercise():
        with (tmp_path / "stdio-stderr.log").open("w") as stderr:
            async with stdio_client(parameters, errlog=stderr) as (read, write):
                async with ClientSession(read, write) as session:
                    await session.initialize()
                    inventory = await session.list_tools()
                    names = {tool.name for tool in inventory.tools}
                    manifest = json.loads((root / "manifest.json").read_text())
                    assert names == {tool["name"] for tool in manifest["tools"]}
                    tool = next(
                        tool
                        for tool in inventory.tools
                        if tool.name == "run_workbook_recipe"
                    )
                    assert tool.outputSchema is None
                    assert (
                        tool.inputSchema["properties"]["mode"]["default"] == "preview"
                    )

                    async def invoke(arguments):
                        response = await session.call_tool(
                            "run_workbook_recipe", arguments
                        )
                        assert not response.isError
                        text = response.content[0].text
                        assert len(text) <= 50000
                        return json.loads(text)

                    arguments = {"filepath": str(source_book), "recipe": recipe}
                    before = source_book.read_bytes()
                    preview = await invoke(arguments)
                    assert preview["ok"] and preview["data"]["ready_to_commit"]
                    assert source_book.read_bytes() == before
                    arguments.update(
                        mode="commit",
                        recipe_token=preview["data"]["recipe_token"],
                        expected_workbook_sha256=preview["data"][
                            "expected_workbook_sha256"
                        ],
                    )
                    result = await invoke(arguments)
                    assert result["ok"] and result["data"]["persisted"]
                    saved = source_book.read_bytes()
                    retry = await invoke(arguments)
                    assert not retry["ok"]
                    assert source_book.read_bytes() == saved

    asyncio.run(asyncio.wait_for(exercise(), timeout=30))
    wb = load_workbook(source_book)
    try:
        assert wb["Report"]["B2"].value == 15
    finally:
        wb.close()


def test_recipe_mcp_file_transport_path_containment(
    source_book, recipe, monkeypatch, tmp_path
):
    from excel_mcp import server

    sandbox = tmp_path / "sandbox"
    sandbox.mkdir()
    monkeypatch.setattr(server, "EXCEL_FILES_PATH", str(sandbox))
    payload = json.loads(server.run_workbook_recipe(str(source_book), recipe))
    assert payload["ok"] is False
    assert "must stay within EXCEL_FILES_PATH" in payload["error"]["message"]


def test_recipe_handles_empty_filtered_output_without_old_rows(source_book, recipe):
    recipe["output"]["allow_overwrite"] = True
    _commit(source_book, recipe, run_workbook_recipe(str(source_book), recipe))
    recipe["filters"] = [{"field": "Amount", "op": "gt", "value": 1000}]
    preview = run_workbook_recipe(str(source_book), recipe)
    assert preview["output_row_count"] == 0
    assert preview["verification"]["passed"]
    _commit(source_book, recipe, preview)
    wb = load_workbook(source_book)
    assert wb["Report"]["A1"].value == "Region"
    assert wb["Report"]["A2"].value is None
    assert wb["Report"]["B3"].value is None
    wb.close()


def test_recipe_preserves_totals_and_unrelated_output_cells(source_book, recipe):
    def add_totals(wb):
        ws = wb["Data"]
        ws.append([None, "Total", "=SUM(C2:C4)"])
        table = ws.tables["Sales"]
        table.ref = "A1:C5"
        table.totalsRowCount = 1
        report = wb.create_sheet("Report")
        report["D1"] = "KEEP"
        report["D1"].number_format = "0.000"

    _edit(source_book, add_totals)
    preview = run_workbook_recipe(str(source_book), recipe)
    assert preview["source_row_count"] == 3
    assert preview["source"]["totals_rows_excluded"] == 1
    _commit(source_book, recipe, preview)
    wb = load_workbook(source_book)
    assert wb["Data"]["C5"].value == "=SUM(C2:C4)"
    assert wb["Report"]["D1"].value == "KEEP"
    assert wb["Report"]["D1"].number_format == "0.000"
    wb.close()


def test_recipe_ignores_formulas_in_unreferenced_columns(source_book, recipe):
    _edit(source_book, lambda wb: setattr(wb["Data"]["A2"], "value", "=1+1"))
    assert run_workbook_recipe(str(source_book), recipe)["ready_to_commit"]
    recipe["checks"] = [{"type": "unique_key", "fields": ["Id"]}]
    with pytest.raises(ValidationError, match="formula"):
        run_workbook_recipe(str(source_book), recipe)


@pytest.mark.parametrize("value", ["20", True])
def test_recipe_rejects_filter_operand_coercion(source_book, recipe, value):
    recipe["filters"][0]["value"] = value
    with pytest.raises(ValidationError, match="operand type"):
        run_workbook_recipe(str(source_book), recipe)


def test_recipe_requires_commit_credentials_and_valid_modes(source_book, recipe):
    with pytest.raises(ValidationError, match="expected_workbook_sha256"):
        run_workbook_recipe(str(source_book), recipe, mode="commit")
    with pytest.raises(ValidationError, match="recipe_token"):
        run_workbook_recipe(
            str(source_book), recipe, mode="commit", expected_workbook_sha256="0" * 64
        )
    for mode in ([], {}, True, "execute"):
        with pytest.raises(ValidationError, match="mode"):
            run_workbook_recipe(str(source_book), recipe, mode=mode)
    with pytest.raises(ValidationError, match="recipe_token"):
        run_workbook_recipe(
            str(source_book), recipe, recipe_token="recipe_v1_" + "\u00e4" * 64
        )


def test_recipe_preserves_source_symlink(source_book, recipe, tmp_path):
    link = tmp_path / "link.xlsx"
    link.symlink_to(source_book)
    preview = run_workbook_recipe(str(link), recipe)
    _commit(link, recipe, preview)
    assert link.is_symlink()
    wb = load_workbook(source_book)
    assert wb["Report"]["B2"].value == 15
    wb.close()


def test_recipe_rejects_chartsheet_output(source_book, recipe):
    from openpyxl.chart import BarChart, Reference

    def chart_sheet(wb):
        chart = BarChart()
        chart.add_data(
            Reference(wb["Data"], min_col=3, min_row=1, max_row=4),
            titles_from_data=True,
        )
        wb.create_chartsheet("Report").add_chart(chart)

    _edit(source_book, chart_sheet)
    with pytest.raises(ValidationError, match="chartsheet"):
        run_workbook_recipe(str(source_book), recipe)


def test_recipe_unicode_response_budgets_match_wire_encoding(
    source_book, recipe, monkeypatch
):
    recipe["metrics"][0]["as"] = "\u00e4" * 128
    preview = run_workbook_recipe(str(source_book), recipe)
    assert preview["ready_to_commit"]
    before = source_book.read_bytes()
    monkeypatch.setattr(recipes.changeset, "MAX_CHANGESET_RESPONSE_BYTES", 1000)
    with pytest.raises(ResponseTooLargeError):
        _commit(source_book, recipe, preview)
    assert source_book.read_bytes() == before


def test_changeset_wire_budget_counts_escaped_unicode(monkeypatch):
    monkeypatch.setattr(recipes.changeset, "MAX_CHANGESET_RESPONSE_BYTES", 400)
    with pytest.raises(ResponseTooLargeError):
        recipes.changeset._assert_response_budget({"value": "\u00e4" * 100})
