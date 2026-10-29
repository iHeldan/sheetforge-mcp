"""Synthetic alternatives, competing editor and complete readback over stdio MCP."""

import asyncio
from decimal import Decimal
import json
from pathlib import Path
import sys

from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table


def test_stdio_scenario_selection_and_global_guard_drift(tmp_path):
    root = Path(__file__).resolve().parents[1]
    path = tmp_path / "planning.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Budget", "Id", "Unit", "Note"])
    for i in range(120):
        ws.append([100, i, f"Unit-{i // 30}", "keep"])
    ws.add_table(Table(displayName="Planning", ref="A1:D121"))
    ws["F1"] = "=1+1"
    wb.save(path)
    wb.close()
    spec = {
        "version": 1,
        "recipe": {
            "version": 1,
            "source": {"table_name": "Planning"},
            "group_by": ["Unit"],
            "metrics": [{"op": "sum", "field": "Budget", "as": "Total"}],
            "output": {"sheet_name": "NotWritten", "range_ref": "A1:B20"},
        },
        "key_fields": ["Id"],
        "guards": [
            {"metric": "Total", "scope": "overall", "max": 12000},
            {"metric": "Total", "scope": "each_group", "min": 2000},
        ],
        "scenarios": [
            {"id": "over", "updates": [{"key": {"Id": 0}, "values": {"Budget": 200}}]},
            {
                "id": "balanced",
                "updates": [
                    {"key": {"Id": i}, "values": {"Budget": 120 if i < 30 else 80}}
                    for i in range(60)
                ],
            },
            {
                "id": "under",
                "updates": [
                    {"key": {"Id": i}, "values": {"Budget": 10}} for i in range(90, 120)
                ],
            },
        ],
    }
    sizes = []
    parameters = StdioServerParameters(
        command=sys.executable,
        args=["-m", "excel_mcp", "stdio"],
        cwd=root,
        env={"PYTHONPATH": str(root / "src")},
    )

    async def exercise():
        with (tmp_path / "stderr.log").open("w") as stderr:
            async with stdio_client(parameters, errlog=stderr) as (read, write):
                async with ClientSession(read, write) as session:
                    await session.initialize()
                    inventory = await session.list_tools()
                    assert len(inventory.tools) == 83
                    tool = next(
                        t
                        for t in inventory.tools
                        if t.name == "evaluate_workbook_scenarios"
                    )
                    assert (
                        tool.annotations.destructiveHint and tool.outputSchema is None
                    )
                    assert set(tool.inputSchema["required"]) == {
                        "filepath",
                        "scenario_set",
                    }
                    args = {
                        "filepath": str(path),
                        "scenario_set": spec,
                        "sample_limit": 1,
                    }

                    async def request(extra=None):
                        reply = await session.call_tool(
                            "evaluate_workbook_scenarios", {**args, **(extra or {})}
                        )
                        wire = "".join(
                            c.text for c in reply.content if c.type == "text"
                        )
                        sizes.append(len(wire))
                        assert len(wire) <= 50000
                        return json.loads(wire)

                    def commit_args(preview, index):
                        return {
                            "mode": "commit",
                            "selected_scenario": preview["scenarios"][index]["id"],
                            "expected_workbook_sha256": preview[
                                "expected_workbook_sha256"
                            ],
                            "scenario_token": preview["scenarios"][index][
                                "scenario_token"
                            ],
                        }

                    original, mtime = path.read_bytes(), path.stat().st_mtime_ns
                    preview = (await request())["data"]
                    assert [s["ready_to_commit"] for s in preview["scenarios"]] == [
                        False,
                        True,
                        False,
                    ]
                    assert (
                        path.read_bytes() == original
                        and path.stat().st_mtime_ns == mtime
                    )
                    denied = await request(commit_args(preview, 0))
                    assert not denied["ok"] and path.read_bytes() == original
                    wb = load_workbook(path)
                    wb["Data"]["D121"] = "human comment"
                    wb.save(path)
                    wb.close()
                    human = path.read_bytes()
                    stale = await request(commit_args(preview, 1))
                    assert not stale["ok"] and path.read_bytes() == human
                    fresh = (await request())["data"]
                    saved = await request(commit_args(fresh, 1))
                    assert saved["ok"] and saved["data"]["persisted"]
                    wb = load_workbook(path)
                    rows = list(
                        wb["Data"].iter_rows(
                            min_row=2, max_row=121, max_col=4, values_only=True
                        )
                    )
                    groups = {}
                    for amount, key, unit, note in rows:
                        assert amount == (120 if key < 30 else 80 if key < 60 else 100)
                        groups[unit] = groups.get(unit, Decimal(0)) + Decimal(
                            str(amount)
                        )
                        assert note == ("human comment" if key == 119 else "keep")
                    assert sum(groups.values()) == Decimal(12000)
                    assert min(groups.values()) == Decimal(2400)
                    assert (
                        wb.sheetnames == ["Data"] and wb["Data"]["F1"].value == "=1+1"
                    )
                    assert wb["Data"].tables["Planning"].ref == "A1:D121"
                    # All proposed cells now satisfy their desired values, but a
                    # different row can invalidate the global outcome.
                    wb["Data"]["A102"] = 1100
                    wb.save(path)
                    wb.close()
                    changed = path.read_bytes()
                    blocked = (await request())["data"]
                    balanced = blocked["scenarios"][1]
                    assert balanced["changed_cell_count"] == 0
                    assert not balanced["ready_to_commit"]
                    assert balanced["metrics"][0]["after"] == "13000"
                    denied = await request(commit_args(blocked, 1))
                    assert not denied["ok"] and path.read_bytes() == changed
                    malformed = await request({"scenario_set": {**spec, "force": True}})
                    assert not malformed["ok"] and path.read_bytes() == changed

    asyncio.run(exercise())
    assert len(sizes) == 8 and sum(sizes) <= 100000
