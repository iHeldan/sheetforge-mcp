from fractions import Fraction
import copy
import hashlib
import json
import random

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table
from openpyxl.utils import get_column_letter

import excel_mcp.explanations as explanations
from excel_mcp.explanations import (
    _breakdown,
    _contract,
    _groups,
    _index,
    _reconcile,
    _text,
    explain_recipe_delta,
)
from excel_mcp.exceptions import (
    PreconditionFailedError,
    ResponseTooLargeError,
    ValidationError,
)


@pytest.fixture
def recipe():
    return {
        "version": 1,
        "source": {"table_name": "Sales"},
        "output": {"sheet_name": "Report", "range_ref": "A1:C20"},
        "filters": [],
        "checks": [],
        "group_by": ["Region"],
        "metrics": [
            {"op": "sum", "field": "Amount", "as": "Revenue"},
            {"op": "count", "as": "Count"},
        ],
    }


def _source(rows):
    return {
        "headers": ["Id", "Region", "Amount"],
        "rows": rows,
        "data_start_row": 2,
        "sheet_name": "Data",
        "table_name": "Sales",
    }


@pytest.fixture
def books(tmp_path):
    def build(name, rows, headers=None):
        path = tmp_path / name
        wb = Workbook()
        wb.active.title = "Data"
        wb.active.append(headers or ["Id", "Region", "Amount"])
        for row in rows:
            wb.active.append(row)
        end = get_column_letter(len(headers or ["Id", "Region", "Amount"]))
        wb.active.add_table(Table(displayName="Sales", ref=f"A1:{end}{len(rows) + 1}"))
        wb.save(path)
        wb.close()
        return path

    return build


def _call(before, after, recipe, **kwargs):
    return explain_recipe_delta(str(before), str(after), recipe, ["Id"], **kwargs)


def test_file_api_explains_all_changes_without_modifying_either_source(books, recipe):
    before = books(
        "before.xlsx", [[1, "North", 60000], [2, "South", 35000], [3, "West", 5000]]
    )
    after = books(
        "after.xlsx", [[2, "South", 32000], [1, "North", 60000], [4, "North", 20000]]
    )
    frozen = [(path.read_bytes(), path.stat().st_mtime_ns) for path in (before, after)]
    result = _call(before, after, recipe)
    revenue, count = result["metrics"]
    assert result["status"] == "verified"
    assert revenue["before"] == "100000" and revenue["after"] == "112000"
    assert revenue["delta"] == "12000" and revenue["residual"] == "0"
    assert revenue["components"] == {
        "added": "20000",
        "removed": "-5000",
        "updated": "-3000",
    }
    assert count["delta"] == "0"
    assert result["verification"]["saved_outputs_checked"] is False
    assert result["verification"]["business_causes_inferred"] is False
    assert result["changes"]["counts"] == {
        "added": 1,
        "removed": 1,
        "updated": 1,
        "unchanged": 1,
    }
    for name, path, (content, mtime) in zip(
        ("before", "after"), (before, after), frozen
    ):
        assert path.read_bytes() == content and path.stat().st_mtime_ns == mtime
        assert result["sources"][name]["sha256"] == hashlib.sha256(content).hexdigest()
    assert {path.name for path in before.parent.iterdir()} == {before.name, after.name}


def test_reordering_is_not_a_data_change(books, recipe):
    before = books("before.xlsx", [[1, "North", 10], [2, "South", 20]])
    after = books(
        "after.xlsx", [[20, 2, "South"], [10, 1, "North"]], ["Amount", "Id", "Region"]
    )
    result = _call(before, after, recipe)
    assert result["changes"]["counts"]["unchanged"] == 2
    assert result["changes"]["sample"] == [] and result["groups"]["changed_count"] == 0
    assert result["sources"]["after"]["bindings"]["Amount"] == "A1"


def test_group_drilldown_and_omitted_cancelling_groups(books, recipe):
    before = books("before.xlsx", [[1, "North", 10]])
    after = books("after.xlsx", [[1, "South", 10]])
    overall = _call(before, after, recipe, sample_limit=1)
    assert overall["metrics"][0]["delta"] == "0"
    assert overall["groups"]["changed_count"] == 2 and overall["groups"]["truncated"]
    shown = overall["groups"]["sample"][0]["metrics"]["Revenue"]["delta"]
    assert int(shown) + int(overall["groups"]["omitted_recipe_delta"]["Revenue"]) == 0
    scoped = _call(before, after, recipe, metric="Revenue", group={"Region": "North"})
    assert scoped["metrics"][0]["delta"] == "-10"
    assert scoped["metrics"][0]["components"]["updated"] == "-10"
    assert scoped["changes"]["sample"][0]["key"] == {"Id": 1}


def test_filter_membership_and_amount_changes_are_not_double_counted(books, recipe):
    recipe["filters"] = [{"field": "Amount", "op": "gt", "value": 10}]
    before = books("before.xlsx", [[1, "North", 5], [2, "South", 20]])
    after = books("after.xlsx", [[1, "North", 15], [2, "South", 5]])
    result = _call(before, after, recipe)
    assert result["metrics"][0]["components"]["updated"] == "-5"
    assert result["metrics"][0]["delta"] == "-5"
    assert result["metrics"][1]["delta"] == "0"
    assert result["changes"]["counts"]["updated"] == 2


def test_pins_bind_source_versions_and_recipe(books, recipe):
    before = books("before.xlsx", [[1, "North", 10]])
    after = books("after.xlsx", [[1, "North", 20]])
    first = _call(before, after, recipe)
    pins = {
        "expected_before_sha256": first["sources"]["before"]["sha256"],
        "expected_after_sha256": first["sources"]["after"]["sha256"],
        "expected_recipe_sha256": first["recipe_sha256"],
    }
    assert (
        _call(before, after, recipe, **pins)["explanation_token"]
        == first["explanation_token"]
    )
    changed = copy.deepcopy(recipe)
    changed["checks"] = [{"type": "row_count", "min": 0, "max": 20}]
    with pytest.raises(PreconditionFailedError) as error:
        _call(before, after, changed, **pins)
    assert error.value.code == "recipe_definition_mismatch"
    pins["expected_after_sha256"] = "0" * 64
    with pytest.raises(PreconditionFailedError) as error:
        _call(before, after, recipe, **pins)
    assert error.value.code == "stale_workbook"


@pytest.mark.parametrize("bad", [None, "", "   ", 1])
def test_bad_keys_fail_closed(books, recipe, bad):
    before = books("before.xlsx", [[1, "North", 10], [bad, "South", 20]])
    after = books("after.xlsx", [[1, "North", 20]])
    with pytest.raises((PreconditionFailedError, ValidationError)):
        _call(before, after, recipe)


def test_key_type_drift_is_not_treated_as_removed_and_added(books, recipe):
    before = books("before.xlsx", [[1, "North", 10]])
    after = books("after.xlsx", [["1", "North", 10]])
    with pytest.raises(PreconditionFailedError) as error:
        _call(before, after, recipe)
    assert error.value.code == "incompatible_field_types"


def test_empty_sources_and_global_aggregate(books, recipe):
    before = books("before.xlsx", [])
    after = books("after.xlsx", [[1, "North", None]])
    recipe["group_by"] = []
    result = _call(before, after, recipe, group={})
    assert result["metrics"][0]["delta"] == "0"
    assert result["metrics"][1]["delta"] == "1"
    assert result["changes"]["sample"][0]["after_row"] == 2


def test_failed_recipe_check_is_an_obstacle(books, recipe):
    before = books("before.xlsx", [[1, "North", 10]])
    after = books("after.xlsx", [[1, "North", 20]])
    recipe["checks"] = [{"type": "row_count", "min": 2, "max": 3}]
    with pytest.raises(PreconditionFailedError) as error:
        _call(before, after, recipe)
    assert error.value.code == "recipe_verification_failed"


def test_scope_is_not_a_mutable_reference_to_caller_input(books, recipe):
    before = books("before.xlsx", [[1, "North", 10]])
    after = books("after.xlsx", [[1, "North", 20]])
    group = {"Region": "North"}
    result = _call(before, after, recipe, group=group)
    group["Region"] = "South"
    assert result["scope"]["group"] == {"Region": "North"}


@pytest.mark.parametrize(
    "op, operand",
    [
        ("eq", {"value": "north"}),
        ("neq", {"value": "north"}),
        ("ne", {"value": "north"}),
        ("contains", {"value": "ort"}),
        ("starts_with", {"value": "no"}),
        ("ends_with", {"value": "th"}),
        ("in", {"values": ["north", "south"]}),
        ("not_in", {"values": ["north"]}),
        ("in", {"value": "north"}),
        ("is_blank", {}),
        ("not_blank", {}),
        ("gt", {"value": "North"}),
        ("gte", {"value": "North"}),
        ("lt", {"value": "South"}),
        ("lte", {"value": "South"}),
    ],
)
def test_independent_filter_semantics_match_recipe_execution(
    books, recipe, op, operand
):
    rows = [[1, "North", 10], [2, "south", 20], [3, None, 5], [4, "   ", 3]]
    before = books("before.xlsx", rows)
    after = books("after.xlsx", list(reversed(rows)))
    recipe["filters"] = [{"field": "Region", "op": op, **operand}]
    result = _call(before, after, recipe)
    assert result["verification"]["passed"]
    assert all(item["delta"] == "0" for item in result["metrics"])


@pytest.mark.parametrize(
    "op",
    ["eq", "neq", "gt", "gte", "lt", "lte", "in", "not_in", "is_blank", "not_blank"],
)
def test_numeric_and_null_filter_semantics(books, recipe, op):
    rows = [[1, "North", 10], [2, "South", None], [3, "South", 20]]
    before = books("before.xlsx", rows)
    after = books("after.xlsx", list(reversed(rows)))
    operand = (
        {}
        if op in {"is_blank", "not_blank"}
        else {"values": [10, None]}
        if op in {"in", "not_in"}
        else {"value": 10}
    )
    recipe["filters"] = [{"field": "Amount", "op": op, **operand}]
    assert _call(before, after, recipe)["verification"]["passed"]


def test_composite_keys_distinguish_repeated_ids(books, recipe):
    before = books("before.xlsx", [[1, "North", 10], [1, "South", 20]])
    after = books("after.xlsx", [[1, "North", 12], [1, "South", 20]])
    result = explain_recipe_delta(str(before), str(after), recipe, ["Id", "Region"])
    assert result["changes"]["counts"]["updated"] == 1
    assert result["changes"]["sample"][0]["key"] == {"Id": 1, "Region": "North"}


def test_decimal_engine_difference_is_explicit_not_a_business_component(books, recipe):
    before = books("before.xlsx", [[1, "North", 0]])
    after = books(
        "after.xlsx", [[1, "North", 1.234567890123456], [2, "North", 0.000000000000001]]
    )
    result = _call(before, after, recipe)
    item = result["metrics"][0]
    from decimal import Decimal

    assert Decimal(item["delta"]) == sum(
        (Decimal(v) for v in item["components"].values()), Decimal(0)
    ) + Decimal(item["numeric_adjustment"])
    assert item["residual"] == "0"
    assert result["numeric_encoding"] == "exact_decimal_strings"


def test_unicode_budget_is_enforced_without_truncating_calculations(books, recipe):
    before = books(
        "before.xlsx", [["\u754c" * 700 + str(i), "\u5317" * 600, i] for i in range(10)]
    )
    after = books(
        "after.xlsx",
        [["\u754c" * 700 + str(i), "\u5357" * 600, i + 10] for i in range(10)],
    )
    with pytest.raises(ResponseTooLargeError):
        _call(before, after, recipe, sample_limit=10)
    result = _call(before, after, recipe, sample_limit=1)
    assert len(json.dumps(result, ensure_ascii=True)) < 50000
    assert result["metrics"][0]["delta"] == "100"
    assert result["changes"]["counts"]["updated"] == 10
    assert result["changes"]["truncated"]
    assert result["changes"]["sample"][0]["key"]["Id"]["truncated"]


@pytest.mark.parametrize(
    "options",
    [
        {"sample_limit": True},
        {"sample_limit": 0},
        {"sample_limit": 11},
        {"metric": []},
        {"metric": "revenue"},
        {"group": []},
        {"group": {"Region": "Missing"}},
        {"group": {"Region": []}},
        {"group": {"unknown": "North"}},
        {"expected_before_sha256": "invalid"},
    ],
)
def test_invalid_scopes_and_pins_fail_cleanly(books, recipe, options):
    before = books("before.xlsx", [[1, "North", 10]])
    after = books("after.xlsx", [[1, "North", 20]])
    with pytest.raises(ValidationError):
        _call(before, after, recipe, **options)


def test_same_file_and_hardlink_aliases_are_rejected(books, recipe, tmp_path):
    import os

    before = books("before.xlsx", [[1, "North", 10]])
    alias = tmp_path / "alias.xlsx"
    os.link(before, alias)
    for after in (before, alias):
        with pytest.raises(ValidationError, match="distinct files"):
            _call(before, after, recipe)


def test_symlinks_to_distinct_versions_are_preserved(books, recipe, tmp_path):
    before = books("before.xlsx", [[1, "North", 10]])
    after = books("after.xlsx", [[1, "North", 20]])
    alias = tmp_path / "alias.xlsx"
    alias.symlink_to(before)
    assert _call(alias, after, recipe)["metrics"][0]["delta"] == "10"
    assert alias.is_symlink()


def test_saved_output_values_are_not_read_as_historical_truth(books, recipe):
    before = books("before.xlsx", [[1, "North", 10]])
    after = books("after.xlsx", [[1, "North", 20]])
    for path in (before, after):
        wb = load_workbook(path)
        wb.create_sheet("Report")["A1"] = "UNRELATED MANUAL OUTPUT"
        wb["Report"].protection.sheet = True
        wb.save(path)
        wb.close()
    result = _call(before, after, recipe)
    assert result["metrics"][0]["delta"] == "10"
    assert result["verification"]["saved_outputs_checked"] is False


def test_randomized_groups_match_independent_decimal_oracle(books, recipe):
    from decimal import Decimal

    rng = random.Random(41927)
    rows = {
        i: [i, rng.choice(["North", "South", None]), rng.randrange(-10000, 10000) / 100]
        for i in range(80)
    }
    before = books("before.xlsx", list(rows.values()))
    changed = {i: list(row) for i, row in rows.items() if i % 9}
    for i, row in changed.items():
        if i % 4 == 0:
            row[2] = rng.randrange(-10000, 10000) / 100
        if i % 7 == 0:
            row[1] = "West"
    for i in range(80, 90):
        changed[i] = [i, "West", rng.randrange(-10000, 10000) / 100]
    after = books("after.xlsx", list(reversed(list(changed.values()))))
    result = _call(before, after, recipe, sample_limit=1)
    expected = sum(
        (Decimal(str(row[2])) for row in changed.values()), Decimal(0)
    ) - sum((Decimal(str(row[2])) for row in rows.values()), Decimal(0))
    revenue = result["metrics"][0]
    assert Decimal(revenue["contribution_delta"]) == expected
    assert (
        Decimal(result["changes"]["sample"][0]["contribution_delta"]["Revenue"])
        + Decimal(result["changes"]["omitted_contribution_delta"]["Revenue"])
        == expected
    )
    for group in {row[1] for row in [*rows.values(), *changed.values()]}:
        detail = _call(before, after, recipe, metric="Revenue", group={"Region": group})
        expected_group = sum(
            (Decimal(str(row[2])) for row in changed.values() if row[1] == group),
            Decimal(0),
        ) - sum(
            (Decimal(str(row[2])) for row in rows.values() if row[1] == group),
            Decimal(0),
        )
        assert Decimal(detail["metrics"][0]["contribution_delta"]) == expected_group


def test_independent_reconciliation_catches_wrong_engine_result(recipe, monkeypatch):
    source = _source([[1, "North", 10]])
    records = _index(source, recipe, ["Id"], ["Id", "Region", "Amount"])
    original = explanations.recipes._aggregate

    def incorrect(*args):
        result = original(*args)
        result["rows"][0][1] += 5
        return result

    monkeypatch.setattr(explanations.recipes, "_aggregate", incorrect)
    with pytest.raises(PreconditionFailedError) as error:
        _reconcile(source, records, _groups(records, recipe), recipe)
    assert error.value.code == "independent_reconciliation_failed"


def test_float_in_another_group_does_not_relax_integer_group_verification(
    recipe, monkeypatch
):
    source = _source([[1, "North", 10**16], [2, "South", 0.1]])
    records = _index(source, recipe, ["Id"], ["Id", "Region", "Amount"])
    original = explanations.recipes._aggregate

    def incorrect(*args):
        result = original(*args)
        result["rows"][0][1] += 1
        return result

    monkeypatch.setattr(explanations.recipes, "_aggregate", incorrect)
    with pytest.raises(PreconditionFailedError, match="independent exact arithmetic"):
        _reconcile(source, records, _groups(records, recipe), recipe)


@pytest.mark.parametrize(
    "old,new", [(1e308, -1e308), (1e-308, -1e-308), (0.1, 0.3), (-3, None)]
)
def test_extreme_values_and_nulls_have_exact_delta_strings(books, recipe, old, new):
    from decimal import Decimal, Context, localcontext

    before = books("before.xlsx", [[1, "North", old]])
    after = books("after.xlsx", [[1, "North", new]])
    with localcontext(Context(prec=1000)):
        expected = Decimal(str(new or 0)) - Decimal(str(old or 0))
    result = _call(before, after, recipe)
    assert Decimal(result["metrics"][0]["contribution_delta"]) == expected
    assert result["metrics"][0]["residual"] == "0"


def test_complete_file_api_is_independent_of_decimal_context(books, recipe):
    from decimal import localcontext

    before = books("before.xlsx", [[1, "North", 1234.5]])
    after = books("after.xlsx", [[1, "North", 1234.6]])
    with localcontext() as context:
        context.prec = 3
        result = _call(before, after, recipe)
    assert result["metrics"][0]["contribution_delta"] == "0.1"


def test_staging_failure_cleans_private_workspaces(books, recipe, monkeypatch):
    from pathlib import Path

    before = books("before.xlsx", [[1, "North", 10]])
    after = books("after.xlsx", [[1, "North", 20]])
    paths = []
    original = explanations.recipes.tempfile.mkdtemp

    def track(*args, **kwargs):
        path = original(*args, **kwargs)
        paths.append(Path(path))
        return path

    def fail(*args):
        raise PreconditionFailedError(
            "source changed during capture", code="stale_workbook"
        )

    monkeypatch.setattr(explanations.recipes.tempfile, "mkdtemp", track)
    monkeypatch.setattr(explanations.changeset, "_stage_workbook", fail)
    with pytest.raises(PreconditionFailedError):
        _call(before, after, recipe)
    assert paths and all(not path.exists() for path in paths)


def test_response_describes_captured_versions_not_later_live_edits(
    books, recipe, monkeypatch
):
    before = books("before.xlsx", [[1, "North", 10]])
    after = books("after.xlsx", [[1, "North", 20]])
    captured_sha = hashlib.sha256(after.read_bytes()).hexdigest()
    original = explanations._reconcile

    def intervene(*args):
        wb = load_workbook(after)
        wb["Data"]["C2"] = 999
        wb.save(after)
        wb.close()
        return original(*args)

    monkeypatch.setattr(explanations, "_reconcile", intervene)
    result = _call(before, after, recipe)
    assert result["metrics"][0]["delta"] == "10"
    assert result["sources"]["after"]["sha256"] == captured_sha
    assert (
        result["sources"]["after"]["sha256"]
        != hashlib.sha256(after.read_bytes()).hexdigest()
    )
    with pytest.raises(PreconditionFailedError, match="version changed"):
        _call(before, after, recipe, expected_after_sha256=captured_sha)


def test_totals_and_unreferenced_formulas_are_excluded(books, recipe):
    before = books(
        "before.xlsx", [[1, "North", 10, "=1+1"]], ["Id", "Region", "Amount", "Memo"]
    )
    after = books(
        "after.xlsx", [[1, "North", 20, "=2+2"]], ["Id", "Region", "Amount", "Memo"]
    )
    for path in (before, after):
        wb = load_workbook(path)
        wb["Data"].append(["TOTAL", None, "=SUM(C2:C2)"])
        wb["Data"].tables["Sales"].ref = "A1:D3"
        wb["Data"].tables["Sales"].totalsRowCount = 1
        wb.save(path)
        wb.close()
    result = _call(before, after, recipe)
    assert result["sources"]["before"]["totals_rows_excluded"] == 1
    assert result["sources"]["before"]["row_count"] == 1
    assert result["metrics"][0]["delta"] == "10"


@pytest.mark.parametrize(
    "field,value", [("Id", "=1+1"), ("Amount", "=SUM(1,2)"), ("Region", "#REF!")]
)
def test_referenced_formulas_and_errors_are_rejected(books, recipe, field, value):
    row = {"Id": 1, "Region": "North", "Amount": 10}
    row[field] = value
    before = books("before.xlsx", [[row[name] for name in row]])
    after = books("after.xlsx", [[1, "North", 20]])
    with pytest.raises(ValidationError, match="formula or Excel error"):
        _call(before, after, recipe)


def test_explanation_and_source_bounds_fail_closed(books, recipe, monkeypatch):
    before = books("before.xlsx", [[1, "North", 10]])
    after = books("after.xlsx", [[1, "North", 20]])
    frozen = [p.read_bytes() for p in (before, after)]
    monkeypatch.setattr(explanations, "MAX_RESPONSE_CHARS", 1)
    with pytest.raises(ResponseTooLargeError):
        _call(before, after, recipe)
    assert [p.read_bytes() for p in (before, after)] == frozen
    monkeypatch.setattr(explanations.recipes, "MAX_SOURCE_CELLS", 1)
    with pytest.raises(ValidationError, match="Source table exceeds"):
        _call(before, after, recipe)


def test_compensated_recipe_and_exact_decimal_comparison_agree(recipe):
    source = _source([[1, "North", 1e16], [2, "North", 1], [3, "North", -1e16]])
    records = _index(source, recipe, ["Id"], ["Id", "Region", "Amount"])
    observed, _ = _reconcile(source, records, _groups(records, recipe), recipe)
    assert next(iter(observed.values())) == [1, 3]


@pytest.mark.parametrize("op", ["avg", "count_distinct", "min", "max"])
def test_contract_rejects_nonadditive_metrics(recipe, op):
    from excel_mcp.exceptions import ValidationError

    recipe.update(
        version=1,
        source={"table_name": "Sales"},
        output={"sheet_name": "Report", "range_ref": "A1:C10"},
    )
    recipe["metrics"][0]["op"] = op
    with pytest.raises(ValidationError, match="sum and count"):
        _contract(recipe, ["Id"])


def test_exact_identity_decomposition(recipe):
    fields = ["Id", "Region", "Amount"]
    before = _index(
        _source([[1, "North", 60000], [2, "South", 35000], [3, "West", 5000]]),
        recipe,
        ["Id"],
        fields,
    )
    after = _index(
        _source([[2, "South", 32000], [1, "North", 60000], [4, "North", 20000]]),
        recipe,
        ["Id"],
        fields,
    )
    result = _breakdown(before, after, fields, 2)
    assert result["counts"] == {"added": 1, "removed": 1, "updated": 1, "unchanged": 1}
    assert result["totals"]["added"] == [20000, 1]
    assert result["totals"]["removed"] == [-5000, -1]
    assert result["totals"]["updated"] == [-3000, 0]
    assert sum(total[0] for total in result["totals"].values()) == 12000


def test_regrouping_cancels_globally_but_not_per_group(recipe):
    fields = ["Id", "Region", "Amount"]
    before = _index(_source([[1, "North", 0.1]]), recipe, ["Id"], fields)
    after = _index(_source([[1, "South", 0.1]]), recipe, ["Id"], fields)
    group = next(iter(_groups(before, recipe)))
    assert _breakdown(before, after, fields, 2)["totals"]["updated"] == [0, 0]
    assert _breakdown(before, after, fields, 2, group)["totals"]["updated"] == [
        -Fraction(1, 10),
        -1,
    ]


@pytest.mark.parametrize(
    "rows, code",
    [
        ([[1, "North", 1], [1, "South", 2]], "duplicate_source_key"),
        ([[None, "North", 1]], "blank_source_key"),
        ([["   ", "North", 1]], "blank_source_key"),
    ],
)
def test_keys_are_validated_even_when_all_rows_are_filtered_out(recipe, rows, code):
    recipe["filters"] = [{"field": "Amount", "op": "gt", "value": 99999}]
    with pytest.raises(PreconditionFailedError) as error:
        _index(_source(rows), recipe, ["Id"], ["Id", "Region", "Amount"])
    assert error.value.code == code


@pytest.mark.parametrize(
    "value", ["0.000000000000000000000001", "1e308", "-1e-308", "0.1", "-12.50"]
)
def test_exact_output_does_not_depend_on_decimal_context(value):
    from decimal import Decimal, localcontext

    with localcontext() as context:
        context.prec = 3
        assert Decimal(_text(Fraction(value))) == Decimal(value)
