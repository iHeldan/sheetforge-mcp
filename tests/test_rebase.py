import copy
import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableFormula

from excel_mcp import changeset, rebase
from excel_mcp.exceptions import (
    DataError,
    PreconditionFailedError,
    ResponseTooLargeError,
    ValidationError,
)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture
def setup_patch(tmp_path):
    def build(name, records, headers=None):
        headers = headers or ["Id", "Name", "Text", "Note"]
        path = tmp_path / name
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(headers)
        for record in records:
            ws.append([record.get(h) for h in headers])
        ws.add_table(Table(displayName="Items", ref=f"A1:D{len(records) + 1}"))
        ws["F1"] = "outside"
        ws["F1"].font = Font(bold=True)
        wb.save(path)
        wb.close()
        return path

    original = [
        {"Id": i, "Name": f"Item {i}", "Text": "old", "Note": "note"} for i in (1, 2, 3)
    ]
    base = build("base.xlsx", original)
    current = build(
        "current.xlsx", list(reversed(original)), ["Text", "Note", "Id", "Name"]
    )
    patch = {
        "version": 1,
        "source": {"table_name": "Items"},
        "key_fields": ["Id"],
        "read_fields": ["Name"],
        "updates": [{"key": {"Id": 1}, "values": {"Text": "new"}}],
    }
    return base, current, patch, build


def call(setup, **kwargs):
    base, current, patch, _ = setup
    return rebase.apply_table_patch(str(base), str(current), patch, sha(base), **kwargs)


def commit(setup, preview):
    return call(
        setup,
        mode="commit",
        rebase_token=preview["rebase_token"],
        expected_workbook_sha256=preview["expected_workbook_sha256"],
    )


def edit(path, cell, value):
    wb = load_workbook(path)
    wb["Data"][cell] = value
    wb.save(path)
    wb.close()


def test_reorders_and_unrelated_edits_survive_verified_commit(setup_patch):
    base, current, _, _ = setup_patch
    edit(current, "B4", "human note")
    frozen_base, frozen_current = base.read_bytes(), current.read_bytes()
    preview = call(setup_patch)
    assert preview["counts"]["apply"] == 1 and preview["ready_to_commit"]
    assert preview["sample"][0]["current_cell"] == "A4"
    assert base.read_bytes() == frozen_base and current.read_bytes() == frozen_current
    result = commit(setup_patch, preview)
    assert result["persisted"] and base.read_bytes() == frozen_base
    wb = load_workbook(current)
    assert wb["Data"]["A4"].value == "new"
    assert wb["Data"]["B4"].value == "human note"
    assert wb["Data"]["F1"].value == "outside" and wb["Data"]["F1"].font.bold
    assert wb["Data"].tables["Items"].ref == "A1:D4"
    wb.close()
    snapshot = Path(result["receipt"]["snapshot"]["path"])
    assert snapshot.read_bytes() == frozen_current


@pytest.mark.parametrize(
    "cell,value,status",
    [("A4", "human", "write_conflict"), ("D4", "renamed", "stale_precondition")],
)
def test_conflicts_block_whole_batch(setup_patch, cell, value, status):
    _, current, patch, _ = setup_patch
    patch["updates"].append({"key": {"Id": 2}, "values": {"Text": "also new"}})
    edit(current, cell, value)
    frozen = current.read_bytes()
    preview = call(setup_patch)
    assert not preview["ready_to_commit"] and preview["counts"][status] == 1
    assert preview["counts"]["apply"] == 1
    with pytest.raises(PreconditionFailedError, match="conflicts"):
        commit(setup_patch, preview)
    assert current.read_bytes() == frozen


def test_changed_premise_wins_over_already_satisfied(setup_patch):
    _, current, _, _ = setup_patch
    edit(current, "A4", "new")
    edit(current, "D4", "new name")
    assert call(setup_patch)["counts"]["stale_precondition"] == 1


def test_already_satisfied_is_read_only_and_stale_retry_is_rejected(setup_patch):
    _, current, _, _ = setup_patch
    edit(current, "A4", "new")
    frozen = current.read_bytes()
    preview = call(setup_patch)
    result = commit(setup_patch, preview)
    assert result["counts"]["already_satisfied"] == 1
    assert result["persisted"] is False and current.read_bytes() == frozen
    assert not list(current.parent.glob("*.snapshot*.xlsx"))
    edit(current, "B2", "another change")
    with pytest.raises(PreconditionFailedError, match="changed"):
        commit(setup_patch, preview)


@pytest.mark.parametrize("which", ["base", "current", "patch", "token"])
def test_preview_credentials_are_bound_to_all_inputs(setup_patch, which):
    base, current, patch, _ = setup_patch
    preview = call(setup_patch)
    expected_base = sha(base)
    if which == "base":
        edit(base, "F1", "new base")
    elif which == "current":
        edit(current, "F1", "human")
    elif which == "patch":
        patch["read_fields"] = []
    else:
        preview["rebase_token"] = "rebase_v1_" + "0" * 64
    frozen = current.read_bytes()
    with pytest.raises(PreconditionFailedError):
        rebase.apply_table_patch(
            str(base),
            str(current),
            patch,
            expected_base,
            mode="commit",
            rebase_token=preview["rebase_token"],
            expected_workbook_sha256=preview["expected_workbook_sha256"],
        )
    assert current.read_bytes() == frozen


def test_repreview_recovers_from_benign_intervening_save(setup_patch):
    preview = call(setup_patch)
    edit(setup_patch[1], "B2", "human")
    with pytest.raises(PreconditionFailedError):
        commit(setup_patch, preview)
    fresh = call(setup_patch)
    assert fresh["rebase_token"] != preview["rebase_token"]
    assert commit(setup_patch, fresh)["persisted"]


def test_missing_target_is_conflict_not_insert(setup_patch):
    base, _, patch, build = setup_patch
    current = build("missing.xlsx", [{"Id": 2, "Name": "Item 2", "Text": "old"}])
    setup = base, current, patch, build
    result = call(setup)
    assert result["counts"]["missing_row"] == 1 and not result["ready_to_commit"]
    patch["updates"][0]["key"]["Id"] = 99
    with pytest.raises(ValidationError, match="pinned base"):
        call(setup)


@pytest.mark.parametrize("value", [None, "", "  ", "=1+1", "#REF!"])
def test_invalid_keys_anywhere_in_table_fail(setup_patch, value):
    edit(setup_patch[1], "C2", value)
    with pytest.raises(ValidationError):
        call(setup_patch)


def test_duplicate_keys_even_unselected_fail(setup_patch):
    edit(setup_patch[1], "C2", 2)
    with pytest.raises(ValidationError, match="duplicate keys"):
        call(setup_patch)


@pytest.mark.parametrize(
    "value",
    [
        "=1+1",
        "#REF!",
        "",
        "a\x01b",
        "x" * 32768,
        12345678901234567,
        0.30000000000000004,
    ],
)
def test_unsafe_or_lossy_outputs_fail_before_write(setup_patch, value):
    setup_patch[2]["updates"][0]["values"]["Text"] = value
    frozen = setup_patch[1].read_bytes()
    with pytest.raises(ValidationError):
        call(setup_patch)
    assert setup_patch[1].read_bytes() == frozen


@pytest.mark.parametrize("value", [None, False, True, 0, 1.25, "plain text"])
def test_literal_output_roundtrips_exactly(setup_patch, value):
    setup_patch[2]["updates"][0]["values"]["Text"] = value
    assert commit(setup_patch, call(setup_patch))["persisted"]
    wb = load_workbook(setup_patch[1])
    actual = wb["Data"]["A4"].value
    assert actual == value and isinstance(actual, bool) == isinstance(value, bool)
    wb.close()


@pytest.mark.parametrize("cell", ["A4", "D4"])
def test_formula_or_error_in_referenced_cells_fails(setup_patch, cell):
    edit(setup_patch[1], cell, "=1+1")
    with pytest.raises(ValidationError, match="literals"):
        call(setup_patch)


def test_unreferenced_formula_is_preserved(setup_patch):
    edit(setup_patch[1], "B4", "=1+1")
    commit(setup_patch, call(setup_patch))
    wb = load_workbook(setup_patch[1])
    assert wb["Data"]["B4"].value == "=1+1"
    wb.close()


@pytest.mark.parametrize("kind", ["protected", "merged", "calculated", "chartsheet"])
def test_unsupported_target_structures_fail_closed(setup_patch, kind):
    path = setup_patch[1]
    wb = load_workbook(path)
    ws = wb["Data"]
    if kind == "protected":
        ws.protection.sheet = True
    elif kind == "merged":
        ws.merge_cells("A4:B4")
    elif kind == "calculated":
        ws.tables["Items"].tableColumns[0].calculatedColumnFormula = TableFormula(
            attr_text="1+1"
        )
    else:
        chart = BarChart()
        chart.add_data(
            Reference(ws, min_col=3, min_row=1, max_row=3), titles_from_data=True
        )
        wb.create_chartsheet("Charts").add_chart(chart)
        setup_patch[2]["source"]["sheet_name"] = "Charts"
    wb.save(path)
    wb.close()
    with pytest.raises((ValidationError, DataError)):
        call(setup_patch)


@pytest.mark.parametrize(
    "mutation",
    [
        "unknown",
        "version",
        "keys",
        "reads",
        "duplicate",
        "write_key",
        "write_read",
        "empty",
    ],
)
def test_closed_contract(setup_patch, mutation):
    patch = setup_patch[2]
    if mutation == "unknown":
        patch["force"] = True
    if mutation == "version":
        patch["version"] = True
    if mutation == "keys":
        patch["key_fields"] = []
    if mutation == "reads":
        del patch["read_fields"]
    if mutation == "duplicate":
        patch["updates"] *= 2
    if mutation == "write_key":
        patch["updates"][0]["values"]["Id"] = 4
    if mutation == "write_read":
        patch["updates"][0]["values"]["Name"] = "new"
    if mutation == "empty":
        patch["updates"] = []
    with pytest.raises(ValidationError):
        call(setup_patch)


def test_alias_and_nonregular_paths_fail(setup_patch, tmp_path):
    base, current, patch, _ = setup_patch
    with pytest.raises(ValidationError, match="distinct"):
        rebase.apply_table_patch(str(base), str(base), patch, sha(base))
    alias = tmp_path / "alias.xlsx"
    alias.hardlink_to(base)
    with pytest.raises(ValidationError, match="distinct"):
        rebase.apply_table_patch(str(alias), str(base), patch, sha(base))
    folder = tmp_path / "directory.xlsx"
    folder.mkdir()
    with pytest.raises(ValidationError, match="regular"):
        rebase.apply_table_patch(str(folder), str(current), patch, sha(base))


def test_response_rejection_precedes_commit(setup_patch, monkeypatch):
    preview = call(setup_patch)
    frozen = setup_patch[1].read_bytes()
    monkeypatch.setattr(rebase, "MAX_SUMMARY_CHARS", 1)
    with pytest.raises(ResponseTooLargeError):
        commit(setup_patch, preview)
    assert setup_patch[1].read_bytes() == frozen


def test_race_after_compile_is_detected_by_changeset(setup_patch, monkeypatch):
    real_apply = changeset.apply_workbook_changeset

    def race(*args, **kwargs):
        edit(setup_patch[1], "B4", "intervening human")
        return real_apply(*args, **kwargs)

    monkeypatch.setattr(changeset, "apply_workbook_changeset", race)
    with pytest.raises(PreconditionFailedError):
        call(setup_patch)
    wb = load_workbook(setup_patch[1])
    assert wb["Data"]["A4"].value == "old"
    assert wb["Data"]["B4"].value == "intervening human"
    wb.close()


def test_sample_size_is_not_commit_authority(setup_patch):
    patch = setup_patch[2]
    patch["updates"] = [
        {"key": {"Id": i}, "values": {"Text": "new"}} for i in (1, 2, 3)
    ]
    preview = call(setup_patch, sample_limit=1)
    bigger = call(setup_patch, sample_limit=3)
    assert preview["rebase_token"] == bigger["rebase_token"]
    assert preview["truncated"] and preview["counts"]["apply"] == 3
    assert commit(setup_patch, preview)["persisted"]


def test_mcp_response_and_both_path_boundaries(setup_patch, monkeypatch, tmp_path):
    import excel_mcp.server as server

    base, current, patch, _ = setup_patch
    reply = json.loads(
        server.apply_table_patch(str(base), str(current), patch, sha(base))
    )
    assert reply["ok"] and reply["data"]["ready_to_commit"]
    original = server.get_excel_path
    seen = []

    def guard(path):
        seen.append(path)
        if path == str(current):
            raise ValidationError("outside allowed root")
        return original(path)

    monkeypatch.setattr(server, "get_excel_path", guard)
    loaded = []
    monkeypatch.setattr(
        server, "apply_table_patch_impl", lambda *a, **kw: loaded.append(a)
    )
    reply = json.loads(
        server.apply_table_patch(str(base), str(current), patch, sha(base))
    )
    assert not reply["ok"] and not loaded and seen == [str(base), str(current)]


def test_numeric_and_bool_identity_are_not_confused(setup_patch):
    base, current, patch, build = setup_patch
    base = build(
        "typed-base.xlsx",
        [{"Id": 1, "Name": "n", "Text": "a"}, {"Id": True, "Name": "b", "Text": "b"}],
    )
    current = build(
        "typed-current.xlsx",
        [{"Id": True, "Name": "b", "Text": "b"}, {"Id": 1.0, "Name": "n", "Text": "a"}],
    )
    patch = copy.deepcopy(patch)
    patch["updates"][0]["values"]["Text"] = "n1"
    assert commit((base, current, patch, build), call((base, current, patch, build)))[
        "persisted"
    ]


def test_composite_keys_and_explicit_empty_read_set(setup_patch):
    patch = setup_patch[2]
    patch["key_fields"] = ["Id", "Name"]
    patch["read_fields"] = []
    patch["updates"][0]["key"]["Name"] = "Item 1"
    assert commit(setup_patch, call(setup_patch))["persisted"]


def test_tokens_cannot_be_reused_on_identical_bytes_at_another_path(
    setup_patch, tmp_path
):
    base, current, patch, build = setup_patch
    preview = call(setup_patch)
    other = tmp_path / "other.xlsx"
    other.write_bytes(current.read_bytes())
    with pytest.raises(PreconditionFailedError, match="token"):
        commit((base, other, patch, build), preview)


def test_sparse_span_limit_has_no_partial_effects(setup_patch):
    base, _, patch, build = setup_patch
    records = [{"Id": i, "Name": "n", "Text": "old"} for i in range(102)]
    base = build("large-base.xlsx", records)
    current = build("large-current.xlsx", records)
    patch["updates"] = [
        {"key": {"Id": i}, "values": {"Text": "new"}} for i in range(0, 102, 2)
    ]
    frozen = current.read_bytes()
    with pytest.raises(ValidationError, match="50 contiguous"):
        call((base, current, patch, build))
    assert current.read_bytes() == frozen


def test_unicode_response_budget_is_escaped(setup_patch):
    patch = setup_patch[2]
    patch["updates"] = [
        {"key": {"Id": i}, "values": {"Text": "\u2603" * 400}} for i in (1, 2, 3)
    ]
    edit(setup_patch[1], "A2", "\u2603" * 400)
    edit(setup_patch[1], "A3", "\u2603" * 400)
    edit(setup_patch[1], "A4", "\u2603" * 400)
    with pytest.raises(ResponseTooLargeError):
        call(setup_patch, sample_limit=3)


def test_rollback_restores_current_not_old_base(setup_patch, monkeypatch):
    from excel_mcp.exceptions import WorkbookError

    base, current, _, _ = setup_patch
    edit(current, "B4", "human note to keep")
    preview = call(setup_patch)
    frozen, frozen_base = current.read_bytes(), base.read_bytes()
    verify = changeset.workbook_module._verify_saved_workbook
    failed = False

    def fail_once(path):
        nonlocal failed
        if Path(path).resolve() == current.resolve() and not failed:
            failed = True
            raise OSError("injected live verification failure")
        return verify(path)

    monkeypatch.setattr(changeset.workbook_module, "_verify_saved_workbook", fail_once)
    with pytest.raises(WorkbookError, match="restored"):
        commit(setup_patch, preview)
    assert (
        failed and current.read_bytes() == frozen and base.read_bytes() == frozen_base
    )


@pytest.mark.parametrize("value", [0, 11, True, [], None])
def test_sample_contract(setup_patch, value):
    with pytest.raises(ValidationError):
        call(setup_patch, sample_limit=value)


@pytest.mark.parametrize("mode", ["force", [], None])
def test_mode_contract(setup_patch, mode):
    with pytest.raises(ValidationError):
        call(setup_patch, mode=mode)


def test_retained_snapshot_cannot_drift_even_during_fresh_preview(setup_patch):
    base, current, patch, _ = setup_patch
    pinned = sha(base)
    edit(base, "F1", "changed base")
    with pytest.raises(PreconditionFailedError):
        rebase.apply_table_patch(str(base), str(current), patch, pinned)


def test_changed_header_set_is_not_guessed(setup_patch):
    path = setup_patch[1]
    wb = load_workbook(path)
    wb["Data"]["B1"] = "Changed"
    wb["Data"].tables["Items"].tableColumns[1].name = "Changed"
    wb.save(path)
    wb.close()
    with pytest.raises(ValidationError, match="header set"):
        call(setup_patch)


def test_fifo_is_rejected_before_open(setup_patch, tmp_path):
    import os

    fifo = tmp_path / "pipe.xlsx"
    os.mkfifo(fifo)
    with pytest.raises(ValidationError, match="regular"):
        rebase.apply_table_patch(
            str(fifo), str(setup_patch[1]), setup_patch[2], sha(setup_patch[0])
        )


@pytest.mark.parametrize(
    "value", [10**300, 10**400, float("inf"), float("nan"), "\ud800"]
)
def test_extreme_output_fails_cleanly(setup_patch, value):
    setup_patch[2]["updates"][0]["values"]["Text"] = value
    with pytest.raises(ValidationError):
        call(setup_patch)


def test_conflicts_are_sampled_before_eligible_cells(setup_patch):
    patch = setup_patch[2]
    patch["updates"] = [
        {"key": {"Id": i}, "values": {"Text": "new"}} for i in (1, 2, 3)
    ]
    edit(setup_patch[1], "A2", "human")
    preview = call(setup_patch, sample_limit=1)
    assert preview["sample"][0]["key"] == {"Id": 3}
    assert preview["sample"][0]["status"] == "write_conflict"
    assert preview["counts"]["apply"] == 2


def test_offset_table_totals_excluded_and_full_preservation(setup_patch):
    base, current, patch, _ = setup_patch
    for path in (base, current):
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        for r, record in enumerate(
            (
                ["Id", "Name", "Text", "Note"],
                [1, "Item 1", "old", "keep"],
                [None, None, "total", None],
            ),
            5,
        ):
            for c, value in enumerate(record, 3):
                ws.cell(r, c, value)
        ws.add_table(Table(displayName="Items", ref="C5:F7", totalsRowCount=1))
        wb.save(path)
        wb.close()
    preview = call(setup_patch)
    assert (
        preview["bindings"]["Id"] == "C5"
        and preview["sample"][0]["current_cell"] == "E6"
    )
    commit(setup_patch, preview)
    wb = load_workbook(current)
    assert wb["Data"]["E6"].value == "new"
    assert wb["Data"]["E7"].value == "total" and wb["Data"]["F6"].value == "keep"
    wb.close()
