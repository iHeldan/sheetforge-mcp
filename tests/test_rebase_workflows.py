"""Synthetic two-editor workflow across the actual stdio MCP boundary."""

import asyncio
import hashlib
import json
from pathlib import Path
import sys

from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table


def test_stdio_rebases_120_updates_preserving_second_editor(tmp_path):
    root = Path(__file__).resolve().parents[1]
    base, current = tmp_path / "base.xlsx", tmp_path / "current.xlsx"
    for path, headers, ids in [
        (base, ["Id", "Name", "Text", "Note"], range(120)),
        (current, ["Text", "Note", "Name", "Id"], reversed(range(120))),
    ]:
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(headers)
        for i in ids:
            record = {"Id": i, "Name": f"Item {i}", "Text": "old", "Note": f"note {i}"}
            ws.append([record[h] for h in headers])
        ws.add_table(Table(displayName="Items", ref="A1:D121"))
        ws["F1"] = "=1+1"
        wb.save(path)
        wb.close()
    frozen_base = base.read_bytes()
    pin = hashlib.sha256(frozen_base).hexdigest()
    patch = {
        "version": 1,
        "source": {"table_name": "Items"},
        "key_fields": ["Id"],
        "read_fields": ["Name"],
        "updates": [
            {"key": {"Id": i}, "values": {"Text": f"Reviewed {i}"}} for i in range(120)
        ],
    }
    sizes = []
    parameters = StdioServerParameters(
        command=sys.executable,
        args=["-m", "excel_mcp", "stdio"],
        cwd=root,
        env={"PYTHONPATH": str(root / "src")},
    )

    async def exercise():
        with (tmp_path / "stderr.log").open("w") as stderr:
            async with stdio_client(parameters, errlog=stderr) as (read, write):
                async with ClientSession(read, write) as session:
                    await session.initialize()
                    inventory = await session.list_tools()
                    tool = next(
                        t for t in inventory.tools if t.name == "apply_table_patch"
                    )
                    assert len(inventory.tools) == 81
                    assert (
                        tool.outputSchema is None and tool.annotations.destructiveHint
                    )
                    assert {
                        "base_filepath",
                        "filepath",
                        "patch",
                        "expected_base_sha256",
                    } <= set(tool.inputSchema["required"])
                    args = {
                        "base_filepath": str(base),
                        "filepath": str(current),
                        "patch": patch,
                        "expected_base_sha256": pin,
                        "sample_limit": 2,
                    }

                    async def request(extra=None):
                        reply = await session.call_tool(
                            "apply_table_patch", {**args, **(extra or {})}
                        )
                        wire = "".join(
                            c.text for c in reply.content if c.type == "text"
                        )
                        sizes.append(len(wire))
                        assert len(wire) <= 50000
                        return json.loads(wire)

                    initial = current.read_bytes()
                    preview = (await request())["data"]
                    assert (
                        preview["counts"]["apply"] == 120
                        and preview["operation_count"] == 1
                    )
                    assert current.read_bytes() == initial
                    # A second editor saves a comment after preview.
                    wb = load_workbook(current)
                    wb["Data"]["B2"] = "human comment"
                    wb.save(current)
                    wb.close()
                    after_human = current.read_bytes()
                    stale = await request(
                        {
                            "mode": "commit",
                            "rebase_token": preview["rebase_token"],
                            "expected_workbook_sha256": preview[
                                "expected_workbook_sha256"
                            ],
                        }
                    )
                    assert not stale["ok"] and current.read_bytes() == after_human
                    fresh = (await request())["data"]
                    done = await request(
                        {
                            "mode": "commit",
                            "rebase_token": fresh["rebase_token"],
                            "expected_workbook_sha256": fresh[
                                "expected_workbook_sha256"
                            ],
                        }
                    )
                    assert done["ok"] and done["data"]["persisted"]
                    wb = load_workbook(current)
                    rows = list(wb["Data"].values)
                    for offset, i in enumerate(reversed(range(120)), 1):
                        assert rows[offset][:4] == (
                            f"Reviewed {i}",
                            "human comment" if offset == 1 else f"note {i}",
                            f"Item {i}",
                            i,
                        )
                    assert wb["Data"]["F1"].value == "=1+1"
                    # Guard drift must block even a plan that would otherwise only be a no-op.
                    wb["Data"]["C2"] = "renamed product"
                    wb.save(current)
                    wb.close()
                    frozen = current.read_bytes()
                    conflict = (await request())["data"]
                    assert conflict["counts"]["stale_precondition"] == 1
                    assert conflict["counts"]["already_satisfied"] == 119
                    assert (
                        not conflict["ready_to_commit"]
                        and current.read_bytes() == frozen
                    )
                    bad = await request({"patch": {**patch, "force": True}})
                    assert not bad["ok"] and current.read_bytes() == frozen

    asyncio.run(exercise())
    assert base.read_bytes() == frozen_base
    assert len(sizes) == 6 and sum(sizes) <= 100000
